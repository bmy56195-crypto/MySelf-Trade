"""Myself Trade - 자가 점검.

인터넷, API 키 없이 프로그램의 핵심 로직이 제대로 도는지 확인한다.
가상 시세로 며칠치 매매를 빠르게 돌려 보고 계산이 맞는지 검사한다.

실행:  python selftest.py
"""
from __future__ import annotations

import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from config import Settings
from dashboard import create_app
from market_data import SyntheticMarket, parse_yahoo_chart
from reports import build_report, export_xlsx
from storage import Storage
from trade_engine import TradeEngine
from utils import calc_trade_pnl, round_step


class Clock:
    def __init__(self, start: datetime):
        self.now = start

    def __call__(self) -> datetime:
        return self.now


class NoAI:
    """점검 중에는 Gemini를 부르지 않는다."""
    available = False


class NoMacro:
    _cache = None

    def fetch(self) -> dict:
        return {"indices": {}, "available": False}


RESULTS: list[tuple[str, bool, str]] = []


def check(name: str, passed: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(passed), detail))
    print(f"  [{'통과' if passed else '실패'}] {name}" + (f" - {detail}" if detail else ""))


def simulate(folder: Path, name: str, days: float, changes: dict, seed: int = 7):
    """가상 시세로 매매를 돌린다. 5분에 한 번 순환하는 것으로 가정."""
    settings = Settings(folder / f"{name}.json")
    settings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul",
                     "ai_interval_minutes": 15, "paper_start_balance": 1000.0, **changes})
    storage = Storage(folder / f"{name}.db")
    start = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)   # 한국 시간 9월 2일 0시
    clock = Clock(start)
    market = SyntheticMarket(seed=seed, history_bars=3000, start_ms=int(start.timestamp() * 1000))
    engine = TradeEngine(settings, storage, market=market, clock=clock, analyzer=NoAI(), macro=NoMacro())
    for _ in range(int(days * 288)):
        engine.step()
        market.advance(1)
        clock.now += timedelta(minutes=5)
    return settings, storage, engine, clock


def is_part(trade: dict) -> bool:
    """분할 매도로 먼저 판 몫을 따로 남긴 기록인지."""
    return '"part_of"' in (trade.get("context") or "")


class FakeAI:
    """Gemini 대신 쓰는 가짜 분석기. 정해 둔 답만 돌려준다 (상호 검증 점검용)."""
    available = True
    knowledge = None

    def __init__(self, action="long", confidence=0.9, p_long=0.9, p_short=0.1):
        self.action, self.confidence, self.p_long, self.p_short = action, confidence, p_long, p_short
        self.seen = []

    def daily_plan(self, macro, snapshots, lessons, last_review):
        from strategy import rule_plan
        return {**rule_plan(snapshots, macro), "source": "gemini"}

    def decide(self, macro, snapshots, plan, lessons, open_trades, state, history=None, rule_views=None, balance=None):
        self.seen.append({"rule_views": rule_views, "balance": balance})
        held = {t["symbol"] for t in open_trades}
        return [{"symbol": symbol, "action": "hold" if symbol in held else self.action, "source": "gemini", "entry": "single",
                 "confidence": self.confidence, "p_long": self.p_long, "p_short": self.p_short, "leverage": 5,
                 "sl_pct": 0.8, "tp_pct": 1.6, "reason": "가짜 AI의 판단"} for symbol in snapshots]

    def review(self, day, trades, pnl_pct, plan, lessons):
        from strategy import rule_review
        return rule_review(day, trades, pnl_pct)

    def digest_material(self, kind, title, text="", url="", image_b64="", mime=""):
        return {"summary": f"{title} 요약", "rules": ["가짜 규칙: RSI 60에서 꺾이면 숏"], "long_conditions": [],
                "short_conditions": ["RSI 60 이상에서 하락 전환"], "cautions": []}


def run_market(folder: Path, name: str, days: float, changes: dict, seed: int = 7, analyzer=None, **market_args):
    """시세의 성격(하락장, 상승장, 박스권)을 정해서 돌린다."""
    settings = Settings(folder / f"{name}.json")
    settings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "ai_interval_minutes": 15,
                     "paper_start_balance": 1000.0, **changes})
    storage = Storage(folder / f"{name}.db")
    start = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
    clock = Clock(start)
    market = SyntheticMarket(seed=seed, history_bars=3000, start_ms=int(start.timestamp() * 1000), **market_args)
    engine = TradeEngine(settings, storage, market=market, clock=clock, analyzer=analyzer or NoAI(), macro=NoMacro())
    for _ in range(int(days * 288)):
        engine.step()
        market.advance(1)
        clock.now += timedelta(minutes=5)
    return settings, storage, engine, clock, market


def section_v9(folder: Path, host: dict, storage, engine, clock, settings) -> None:
    import base64
    import json as _json
    import os

    import config as _config
    import gemini_analyzer as _gemini
    import knowledge as _knowledge
    import market_data as _market
    from config import SYMBOLS, SettingsError
    from reports import build_pnl_board
    from risk_manager import RiskManager, daily_limits
    from strategy import rule_scores

    print("\n8. 익절/손절 구조 (전체 원금 기준, 진입 투자금 기준)와 보유 시간")
    v = Settings(folder / "v9.json")
    errors = {}
    for bad in (4, 0, -30, 1441, 99999, 4.9, "abc"):
        try:
            v.update({"max_hold_minutes": bad})
            errors[bad] = ""
        except SettingsError as exc:
            errors[bad] = str(exc)
    check("보유 시간 5분 미만, 1440분 초과, 숫자가 아닌 값은 저장되지 않고 안내 문구가 나옴",
          all(errors.values()) and "최소 5분" in errors[4] and "1440분" in errors[1441] and v.get("max_hold_minutes") == 240, errors[4])
    v.update({"max_hold_minutes": 5})
    low_ok = v.get("max_hold_minutes") == 5
    v.update({"max_hold_minutes": 1440})
    check("보유 시간 5분과 1440분(경계값)은 저장됨", low_ok and v.get("max_hold_minutes") == 1440)
    app = create_app(engine, settings, storage).test_client()
    rejected = app.post("/api/settings", json={"max_hold_minutes": 3}, headers=host)
    check("화면에서 보낸 보유 시간 3분은 서버가 거부", rejected.status_code == 400 and "보유 시간" in rejected.get_json()["error"],
          rejected.get_json()["error"])

    v.update({"daily_target_unit": "amount", "daily_target_amount": 300, "daily_loss_unit": "amount", "daily_loss_amount": 250})
    target_pct, loss_pct = daily_limits(v.all(), 1000.0)
    check("원금 기준 목표/손실을 금액으로 정하면 오늘 원금 대비 %로 바꿔 적용", abs(target_pct - 30) < 1e-9 and abs(loss_pct - 25) < 1e-9,
          f"원금 1000: +300 USDT = +{target_pct:g}%, -250 USDT = -{loss_pct:g}%")
    target2, loss2 = daily_limits(v.all(), 2000.0)
    check("원금이 달라지면 같은 금액의 %도 달라짐", abs(target2 - 15) < 1e-9 and abs(loss2 - 12.5) < 1e-9, f"원금 2000: +{target2:g}%, -{loss2:g}%")
    rm = RiskManager(v)
    check("금액 기준 손실 한도로 잠금 판단", rm.should_lock(-25.1, loss_pct) and not rm.should_lock(-24.9, loss_pct))

    v.update({"tpsl_mode": "manual", "pos_tp_unit": "pct", "pos_tp_pct": 50, "pos_sl_unit": "pct", "pos_sl_pct": 20})
    fixed = rm.manual_sl_tp(10, 0.00055, 100.0)
    tp_net = calc_trade_pnl("long", 100.0, 100.0 * (1 + fixed["tp_pct"] / 100), 10.0, 0.00055, 0.00055)["net_pnl"]
    sl_net = calc_trade_pnl("long", 100.0, 100.0 * (1 - fixed["sl_pct"] / 100), 10.0, 0.00055, 0.00055)["net_pnl"]
    check("투자금 기준 익절 50% / 손절 20%: 수수료를 뺀 순손익이 투자금 100의 +50, -20", abs(tp_net - 50) < 0.1 and abs(sl_net + 20) < 0.1,
          f"10배, 가격 +{fixed['tp_pct']}% / -{fixed['sl_pct']}% -> 순손익 {tp_net:+.2f} / {sl_net:+.2f} USDT")
    short_tp = calc_trade_pnl("short", 100.0, 100.0 * (1 - fixed["tp_pct"] / 100), 10.0, 0.00055, 0.00055)["net_pnl"]
    check("숏도 같은 기준으로 계산", abs(short_tp - 50) < 0.1, f"숏 익절 순손익 {short_tp:+.2f} USDT")
    v.update({"pos_tp_unit": "amount", "pos_tp_amount": 30, "pos_sl_unit": "amount", "pos_sl_amount": 15})
    by_amount = rm.manual_sl_tp(10, 0.00055, 100.0)
    check("금액으로 정하면 포지션 투자금으로 나눠 비율로 바꿈", by_amount["tp_margin_pct"] == 30 and by_amount["sl_margin_pct"] == 15,
          f"투자금 100: +30 USDT = +{by_amount['tp_margin_pct']:g}%, -15 USDT = -{by_amount['sl_margin_pct']:g}%")
    v.update({"pos_sl_amount": 80})
    capped = rm.manual_sl_tp(10, 0.00055, 100.0)
    check("손절은 투자금의 60%까지만 (강제 청산보다 먼저)", capped["sl_margin_pct"] == 60 and "60%" in capped["note"])
    try:
        v.update({"pos_sl_unit": "pct", "pos_sl_pct": 75})
        too_big = False
    except SettingsError:
        too_big = True
    check("손절 비율 60% 초과는 저장되지 않음", too_big)
    from risk_manager import max_sl_move
    v.update({"pos_tp_unit": "pct", "pos_tp_pct": 300, "pos_sl_unit": "pct", "pos_sl_pct": 20})
    far = rm.manual_sl_tp(3, 0.00055, 100.0, 50)
    short_tp_price = rm.sl_tp_prices("short", 100.0, far["sl_pct"], far["tp_pct"], 0.01)[1]
    check("익절 목표가 너무 멀면 가격 50% 움직임까지만 걸어 숏의 익절 가격이 0 아래로 내려가지 않음",
          far["tp_pct"] == 50.0 and short_tp_price > 0 and "50%" in far["note"], f"3배, 익절 300% -> 가격 폭 {far['tp_pct']}%, 숏 익절가 {short_tp_price}")
    v.update({"pos_tp_pct": 50, "pos_sl_pct": 3})
    ceiling = rm.max_leverage_for_stop(3, 0.00055)
    tight = rm.manual_sl_tp(ceiling, 0.00055, 100.0, 50)
    tight_net = calc_trade_pnl("long", 100.0, 100.0 * (1 - tight["sl_pct"] / 100), ceiling, 0.00055, 0.00055)["net_pnl"]
    check("손절 비율이 작으면 그 손실을 지킬 수 있는 배율까지만 씀 (수수료가 손절 금액을 넘지 않게)",
          ceiling < 50 and tight["sl_pct"] >= 0.1 and abs(tight_net + 3) < 0.05,
          f"손절 3%는 {ceiling}배 이하에서만 가능, 그 배율에서 순손실 {tight_net:+.2f} USDT (투자금 100)")
    check("최대 레버리지에 가까울수록 손절을 더 좁게 (강제 청산 거리 안쪽)", max_sl_move(50, 50) == 0.8 and max_sl_move(25, 25) == 1.6
          and max_sl_move(10, 100) == 6.0 and rm.sl_tp_pct(3.0, None, 0.9, 25, 0.00055, 25)[0] == 1.6,
          f"50배(최대 50배) {max_sl_move(50, 50)}%, 25배(최대 25배) {max_sl_move(25, 25)}%, 10배(최대 100배) {max_sl_move(10, 100)}%")
    lsettings, lstorage, _, _, _ = run_market(folder, "tight_sl", 1, {
        "tpsl_mode": "manual", "pos_tp_pct": 20, "pos_sl_pct": 3, "leverage_mode": "manual", "manual_leverage": 50, "entry_mode": "single"})
    lrows = lstorage.closed_trades("paper") + lstorage.open_trades("paper")
    lnotes = [r["note"] or "" for r in lstorage.recent_analysis("paper", 300, executed_only=True)]
    check("50배 고정 + 손절 3%: 거래할 때 배율을 낮추고 이유를 남김", bool(lrows) and all(t["leverage"] == ceiling for t in lrows)
          and any("낮췄습니다" in n for n in lnotes), f"{len(lrows)}건 모두 {ceiling}배")

    msettings, mstorage, mengine, mclock, _ = run_market(folder, "manual_tpsl", 2, {
        "tpsl_mode": "manual", "pos_tp_pct": 30, "pos_sl_pct": 15, "leverage_mode": "manual", "manual_leverage": 10,
        "entry_mode": "single", "max_hold_minutes": 1440})
    done = mstorage.closed_trades("paper")
    wins = [t for t in done if t["exit_reason"] == "익절"]
    cuts = [t for t in done if t["exit_reason"] == "손절"]
    check("수동 익절/손절로 청산된 거래가 있음", bool(wins) and bool(cuts), f"익절 {len(wins)}건, 손절 {len(cuts)}건")
    if wins and cuts:
        check("익절 거래의 순수익이 투자금의 약 +30% 이상", all(t["return_pct"] >= 28.5 for t in wins),
              "수익률 " + ", ".join(f"{t['return_pct']:+.1f}%" for t in wins[:5]) + " (5분마다 가격을 확인하는 가상 시세라 조금 더 간 뒤 체결됨)")
        check("손절 거래의 순손실이 투자금의 약 -15% 근처", all(-30 < t["return_pct"] <= -13.5 for t in cuts),
              "수익률 " + ", ".join(f"{t['return_pct']:+.1f}%" for t in cuts[:5]))
    contexts = [_json.loads(t["context"]).get("tpsl", {}) for t in done]
    check("거래 기록에 익절/손절 방식(수동, 투자금 기준 값)이 남음", bool(contexts) and all(c.get("mode") == "manual" and c.get("tp_margin_pct") == 30 for c in contexts))
    mstate = mengine.get_state()
    check("상태에 익절/손절 방식과 보유 시간 제한 표시", mstate["tpsl_mode"] == "manual" and "+30%" in mstate["tpsl_text"] and mstate["max_hold_minutes"] == 1440,
          mstate["tpsl_text"])
    if mstate["positions"]:
        pos = mstate["positions"][0]
        check("보유 포지션에 익절/손절 시 순손익 표시", abs(pos["tp_margin_pct"] - 30) < 1.0 and abs(pos["sl_margin_pct"] + 15) < 1.0,
              f"익절 시 {pos['tp_net']:+.2f} USDT ({pos['tp_margin_pct']:+.1f}%), 손절 시 {pos['sl_net']:+.2f} USDT ({pos['sl_margin_pct']:+.1f}%)")
    osettings, ostorage, oengine, oclock, _ = run_market(folder, "overnight", 1.6, {"max_hold_minutes": 1440, "entry_mode": "single"}, seed=4)
    auto_rows = ostorage.closed_trades("paper")
    auto_ctx = [_json.loads(t["context"]) for t in auto_rows]
    widths = sorted({round(c["sl_pct"], 2) for c in auto_ctx})
    check("자동 모드에서는 변동성(ATR) 기준으로 거래마다 익절/손절 폭이 달라짐", bool(auto_ctx) and all(c["tpsl"]["mode"] == "auto" for c in auto_ctx)
          and len(widths) > 3, f"손절 폭 {widths[0]}% ~ {widths[-1]}% ({len(widths)}가지)")

    hsettings, hstorage, _, _, _ = run_market(folder, "hold5", 1, {"max_hold_minutes": 5, "entry_mode": "single"})
    held = [(datetime.fromisoformat(t["exit_time"]) - datetime.fromisoformat(t["entry_time"])).total_seconds() / 60 for t in hstorage.closed_trades("paper")]
    check("보유 시간 5분이면 5분 뒤 정리", bool(held) and max(held) <= 5.01, f"{len(held)}건, 가장 긴 보유 {max(held):.0f}분")

    print("\n9. 롱/숏 균형과 AI 상호 검증")

    def mirror(snap: dict) -> dict:
        """시세를 위아래로 뒤집은 상황 (상승 <-> 하락). 롱 점수와 숏 점수가 서로 바뀌어야 한다."""
        flip = {"up": "down", "down": "up"}
        out = {"ok": True}
        center = snap["m15"]["ema20"]          # 모든 가격을 같은 기준선에 대해 뒤집는다
        for frame in ("m15", "h1", "d1"):
            f = dict(snap[frame])
            if not f.get("ok"):
                out[frame] = f
                continue
            f.update(price=2 * center - f["price"], ema20=2 * center - f["ema20"], ema50=2 * center - f["ema50"],
                     trend=flip.get(f["trend"], f["trend"]), regime=flip.get(f["regime"], f["regime"]),
                     structure=flip.get(f["structure"], f["structure"]), rsi=round(100 - f["rsi"], 1), macd_hist=-f["macd_hist"],
                     macd_rising=not f["macd_rising"], bb_pos=round(1 - f["bb_pos"], 2), chg_4=-f["chg_4"], slope_pct=-f["slope_pct"],
                     high_20=2 * center - f["low_20"], low_20=2 * center - f["high_20"])
            out[frame] = f
        return out

    probe = SyntheticMarket(seed=21, history_bars=3000, start_ms=1_790_000_000_000)
    same, total = 0, 0
    for _ in range(60):
        probe.advance(7)
        for symbol in probe.BASE:
            snap = _market.coin_snapshot(probe, symbol)
            if not snap.get("ok") or snap["m15"]["macd_hist"] == 0 or snap["m15"]["chg_4"] == 0:
                continue
            a, b = rule_scores(snap, {}), rule_scores(mirror(snap), {})
            total += 1
            same += abs(a["long"] - b["short"]) < 1e-9 and abs(a["short"] - b["long"]) < 1e-9
    check("롱과 숏을 같은 기준으로 채점 (시세를 뒤집으면 점수도 정확히 뒤바뀜)", total > 100 and same == total, f"{total}개 상황 중 {same}개 일치")

    def sides(st) -> dict:
        rows = st.closed_trades("paper") + st.open_trades("paper")
        return {"long": sum(1 for t in rows if t["side"] == "long"), "short": sum(1 for t in rows if t["side"] == "short")}

    _, dstorage, _, _, _ = run_market(folder, "bear", 1.5, {}, seed=2, trend=-0.00025)
    down = sides(dstorage)
    check("하락장에서는 숏으로 진입", down["short"] >= 5 and down["long"] == 0, f"롱 {down['long']}회, 숏 {down['short']}회")
    _, ustorage, _, _, _ = run_market(folder, "bull", 1.5, {}, seed=2, trend=0.00025)
    up = sides(ustorage)
    check("상승장에서는 롱으로 진입", up["long"] >= 5 and up["short"] == 0, f"롱 {up['long']}회, 숏 {up['short']}회")
    _, bstorage, _, _, _ = run_market(folder, "box", 2, {}, seed=1, revert=0.02)
    box = sides(bstorage)
    box_rows = bstorage.closed_trades("paper") + bstorage.open_trades("paper")
    setups = [_json.loads(t["context"]) for t in box_rows]
    check("박스권에서는 상단 숏과 하단 롱을 함께 씀", box["long"] >= 3 and box["short"] >= 3, f"롱 {box['long']}회, 숏 {box['short']}회")
    in_box = [c for c in setups if c.get("regime") == "range"]
    check("박스권으로 판단한 구간에서는 추세 추종 진입을 하지 않음", bool(in_box) and all(c.get("setup") == "range" for c in in_box),
          f"박스권 진입 {len(in_box)}건 모두 박스 상단/하단 진입")
    box_short = [(t, c) for t, c in zip(box_rows, setups) if c.get("setup") == "range" and t["side"] == "short"]
    check("박스권 숏은 볼린저 상단에서만", bool(box_short) and all(c["bb_pos"] >= 0.85 and c["rsi"] >= 60 for _, c in box_short))

    ai_long = FakeAI("long", 0.9, 0.9, 0.1)
    _, cstorage, _, _, _ = run_market(folder, "cross_on", 1, {"cross_check": True}, seed=2, analyzer=ai_long, trend=-0.00025)
    crossed = sides(cstorage)
    notes = [r["note"] or "" for r in cstorage.recent_analysis("paper", 400)]
    passed = cstorage.closed_trades("paper") + cstorage.open_trades("paper")
    blocked_n = sum("상호 검증" in n for n in notes)
    check("AI에게 규칙 점수와 최근 진입 방향을 함께 전달", bool(ai_long.seen) and "BTCUSDT" in ai_long.seen[-1]["rule_views"]
          and set(ai_long.seen[-1]["balance"]) == {"long", "short"})
    _, fstorage, _, _, _ = run_market(folder, "cross_off", 1, {"cross_check": False}, seed=2, analyzer=FakeAI("long", 0.9, 0.9, 0.1),
                                      trend=-0.00025)
    alone = sides(fstorage)["long"]
    check("하락장에서 AI가 롱만 외쳐도, 규칙 점수가 같은 방향이 아니면 진입하지 않음 (상호 검증)",
          blocked_n > 0 and crossed["long"] * 3 <= alone and all("상호 검증 통과" in (t["entry_reason"] or "") for t in passed),
          f"AI 혼자면 롱 {alone}회 -> 상호 검증을 켜면 {crossed['long']}회 (규칙도 롱을 가리킨 반등 구간만), 막힌 판단 {blocked_n}건")
    ai_short = FakeAI("short", 0.9, 0.1, 0.9)
    _, gstorage, _, _, _ = run_market(folder, "cross_agree", 1, {"cross_check": True}, seed=2, analyzer=ai_short, trend=-0.00025)
    agreed = gstorage.closed_trades("paper") + gstorage.open_trades("paper")
    check("AI와 규칙이 같은 방향이면 진입하고 종합 확신도를 기록", bool(agreed) and all(t["side"] == "short" for t in agreed)
          and all("상호 검증 통과" in (t["entry_reason"] or "") for t in agreed), f"숏 {len(agreed)}회")
    held_long = {"side": "long"}
    check("AI의 추가 매수(분할 다음 회차)도 확신도와 상호 검증을 통과해야 함",
          "확신도" in TradeEngine._add_blocked({"source": "gemini", "confidence": 0.05}, held_long, {"long": 0.7, "short": 0.2}, {"cross_check": True})
          and "상호 검증" in TradeEngine._add_blocked({"source": "gemini", "confidence": 0.8}, held_long, {"long": 0.2, "short": 0.7}, {"cross_check": True})
          and TradeEngine._add_blocked({"source": "gemini", "confidence": 0.8}, held_long, {"long": 0.7, "short": 0.2}, {"cross_check": True}) == "")
    _, vstorage, _, _, _ = run_market(folder, "vague", 0.5, {"cross_check": False}, seed=2, analyzer=FakeAI("long", 0.9, 0.6, 0.55))
    vague = [r["note"] or "" for r in vstorage.recent_analysis("paper", 200)]
    check("AI가 본 롱 확률과 숏 확률의 차이가 작으면 진입하지 않음", not (vstorage.closed_trades("paper") + vstorage.open_trades("paper"))
          and any("차이가" in n for n in vague))

    print("\n10. 손익 현황, 거래 기록 (진입 시각 / 청산 시각, 더보기)")
    bsettings, bstorage2, bengine, bclock, _ = run_market(folder, "board", 0.8, {"max_hold_minutes": 120, "entry_mode": "single", "invest_pct": 40})
    state = bengine.get_state()
    board = build_pnl_board(bstorage2, "paper", "Asia/Seoul", state["positions"], state["start_equity"], now=bclock.now)
    check("코인별 오늘 손익의 합계 = 메인 화면의 오늘 수익", board["active"] and abs(board["net_total"] - state["pnl_amount"]) < 1e-6,
          f"{board['net_total']:+.4f} = {state['pnl_amount']:+.4f} USDT (수익 코인 {len(board['profits'])}개, 손실 코인 {len(board['losses'])}개)")
    check("수익 표와 손실 표로 나뉨", all(c["net"] > 0 for c in board["profits"]) and all(c["net"] < 0 for c in board["losses"])
          and abs(board["profit_total"] + board["loss_total"] - board["net_total"]) < 1e-9)
    check("표에 코인명, 거래 금액, 수익률, 수익금이 들어 있음", all({"name", "notional", "return_pct", "net"} <= set(c) for c in board["coins"]))
    ostate = oengine.get_state()
    oboard = build_pnl_board(ostorage, "paper", "Asia/Seoul", ostate["positions"], ostate["start_equity"], now=oclock.now)
    carried = ostorage.kv_get(f"carry_paper_{ostate['day']}") or {}
    check("날짜를 넘겨 들고 온 포지션이 있어도 합계가 오늘 수익과 같음", bool(carried) and abs(oboard["net_total"] - ostate["pnl_amount"]) < 1e-6,
          f"넘어온 포지션 {len(carried)}개, {oboard['net_total']:+.4f} = {ostate['pnl_amount']:+.4f} USDT")
    storage, clock = ostorage, oclock           # 아래 거래 기록 점검은 이 가상 매매 결과로 한다
    app = create_app(oengine, osettings, ostorage).test_client()
    # 프로그램이 꺼져 있던 사이 날짜가 바뀌고 보유 시간도 지난 경우: 켜자마자 정리되는 거래는 오늘 손익에 0으로 잡혀야 한다
    zsettings, zstorage, zengine, zclock, zmarket = run_market(folder, "offnight", 0.3, {"max_hold_minutes": 60, "entry_mode": "single"}, seed=4)
    for _ in range(288):
        if zstorage.open_trades("paper"):
            break
        zengine.step(); zmarket.advance(1); zclock.now += timedelta(minutes=5)
    carried_over = len(zstorage.open_trades("paper"))
    zsettings.update({"trading_enabled": False})
    zmarket.advance(240); zclock.now += timedelta(hours=20)
    zengine.step()
    zstate = zengine.get_state()
    zboard = build_pnl_board(zstorage, "paper", "Asia/Seoul", zstate["positions"], zstate["start_equity"], now=zclock.now)
    zdays = zstorage.days_between("paper", "2026-01-01", "2027-01-01")
    check("꺼져 있다가 다음 날 켰을 때 바로 정리된 거래는 오늘 손익에 넣지 않음 (이미 오늘 원금에 반영)",
          carried_over > 0 and not zstate["positions"] and abs(zstate["pnl_amount"]) < 1e-6 and abs(zboard["net_total"]) < 1e-6
          and len(zdays) == 2 and abs(zdays[0]["end_equity"] - zdays[1]["start_equity"]) < 1e-6,
          f"넘어온 포지션 {carried_over}개 정리, 오늘 수익 {zstate['pnl_amount']:+.4f}, 손익 현황 합계 {zboard['net_total']:+.4f}")
    total = len(storage.closed_trades("paper"))
    first = app.get("/api/trades?limit=10", headers=host).get_json()
    more = app.get("/api/trades?limit=20", headers=host).get_json()
    check("거래 기록 더보기: 10건 다음에 10건을 이어서 조회", first["ok"] and len(first["trades"]) == min(10, total)
          and len(more["trades"]) == min(20, total) and more["trades"][:10] == first["trades"] and first["has_more"] == (total > 10),
          f"전체 {total}건 중 {len(first['trades'])}건 -> {len(more['trades'])}건")
    everything = app.get(f"/api/trades?limit={total + 5}", headers=host).get_json()
    times = [t["exit_at"] for t in everything["trades"]]
    check("끝까지 조회하면 더보기가 사라짐 (최신순 정렬)", len(times) == total and everything["has_more"] is False and times == sorted(times, reverse=True))
    row = first["trades"][0]
    check("진입 시각과 청산 시각을 따로, 초 단위까지 표시", len(row["entry_at"]) == 19 and len(row["exit_at"]) == 19 and row["entry_at"] < row["exit_at"],
          f"진입 {row['entry_at']} / 청산 {row['exit_at']}")
    sheet = load_workbook(export_xlsx(storage, "paper", "1m", "Asia/Seoul", now=clock.now))["거래내역"]
    head = [c.value for c in sheet[1]]
    check("엑셀에도 진입 시각, 청산 시각 열이 따로 있음", "진입 시각" in head and "청산 시각" in head
          and len(str(sheet.cell(row=2, column=head.index("진입 시각") + 1).value)) == 19)

    ignored = app.post("/api/settings", json={"custom_symbols": [{"symbol": "FAKEUSDT"}]}, headers=host)
    not_object = app.post("/api/settings", json=[1, 2], headers=host)
    check("코인 목록은 설정 저장으로 바꿀 수 없고(Bybit 확인을 거쳐야 함), 형식이 틀린 요청은 거부",
          ignored.status_code == 200 and "FAKEUSDT" not in SYMBOLS and not_object.status_code == 400)

    print("\n11. 글로벌 지수, 금 시세, 환율")
    live_bar = _market.parse_yahoo_chart({"chart": {"result": [{
        "meta": {"regularMarketPrice": 5123.41, "regularMarketTime": 1790000000 + 3 * 86400 + 3600, "gmtoffset": -14400},
        "timestamp": [1790000000 + i * 86400 for i in range(4)],
        "indicators": {"quote": [{"close": [5000.0, 5050.5, 5063.2001953125, 5123.41015625]}]}}]}})
    check("Yahoo의 전일 대비: 종가의 소수 자리가 지금 가격과 조금 달라도 오늘 봉으로 알아봄", live_bar is not None and live_bar["chg_1d"] == 1.19,
          f"전일 5063.20 -> 지금 5123.41 = {live_bar['chg_1d']:+.2f}%")
    closed_bar = _market.parse_yahoo_chart({"chart": {"result": [{
        "meta": {"regularMarketPrice": 5200.0, "regularMarketTime": 1790000000 + 4 * 86400 + 3600, "gmtoffset": -14400},
        "timestamp": [1790000000 + i * 86400 for i in range(4)],
        "indicators": {"quote": [{"close": [5000.0, 5050.5, 5063.2, 5123.41]}]}}]}})
    check("오늘 봉이 아직 없으면 마지막 종가가 전일 종가", closed_bar["chg_1d"] == round((5200.0 / 5123.41 - 1) * 100, 2))
    fred = _market.parse_fred_csv("observation_date,SP500\n2026-09-29,6600.0\n2026-09-30,.\n2026-10-01,6633.0\n")
    check("FRED 자료 해석 (값이 없는 날은 건너뜀)", fred is not None and fred["last"] == 6633.0 and fred["chg_1d"] == 0.5, str(fred))

    class Reply:
        def __init__(self, code=200, data=None, text=""):
            self.status_code, self._data, self.text = code, data, text

        def json(self):
            if self._data is None:
                raise ValueError("no json")
            return self._data

    calls = []

    def fake_get(url, params=None, timeout=8.0):
        calls.append(url)
        if "yahoo" in url:
            if "GSPC" in url or "KRW" in url:
                return Reply(429, text="Too Many Requests")      # Yahoo가 막힌 상황
            return Reply(200, {"chart": {"result": [{"meta": {"regularMarketPrice": 110.0, "regularMarketTime": 1790000000},
                                                     "indicators": {"quote": [{"close": [100.0, 104.0, 108.0, 110.0]}]}}]}})
        if "fred" in url:
            return Reply(200, text="observation_date,SP500\n2026-09-30,6600.0\n2026-10-01,6666.0\n")
        if "upbit" in url:
            return Reply(200, [{"market": "KRW-USDT", "trade_price": 1354.0}])
        return Reply(500)

    real_get = _market.MacroData._get
    _market.MacroData._get = staticmethod(fake_get)
    try:
        macro = _market.MacroData().fetch()
    finally:
        _market.MacroData._get = real_get
    names = [item["name"] for item in macro["items"]]
    check("나스닥, 다우존스, S&P 500, 국제 금 시세 네 가지를 모두 받아 옴", names == ["나스닥", "다우존스", "S&P 500", "국제 금 시세"]
          and all(item.get("last") for item in macro["items"]), ", ".join(f"{i['name']} {i['last']}" for i in macro["items"]))
    check("Yahoo가 막힌 지수는 다른 출처(FRED)에서 대신 받음", macro["indices"]["S&P 500"]["source"].startswith("FRED") and macro["indices"]["S&P 500"]["chg_1d"] == 1.0)
    check("원화 환산용 환율", macro["fx"] == {"rate": 1354.0, "source": "업비트 USDT/KRW"})

    class SlowMacro(NoMacro):
        def __init__(self, data):
            self._cache = data

        def peek(self):
            return self._cache

        def refresh_async(self, force=False):
            return False

    oengine.macro = SlowMacro(macro)
    shown = app.get("/api/macro", headers=host).get_json()
    oengine.macro = NoMacro()
    check("분석 화면용 글로벌 지수 조회 (Gemini 키가 없어도 표시)", shown["ok"] and len(shown["macro"]["items"]) == 4 and shown["gemini"]["state"] in ("none", "unchecked", "ok", "error"))
    empty = app.get("/api/macro", headers=host).get_json()
    check("지수 데이터가 아직 없어도 화면이 깨지지 않음", empty["ok"] and empty["macro"] is None)

    print("\n12. Gemini 연결 상태와 학습 자료")
    saved_key = os.environ.get("GEMINI_API_KEY")
    real_post = _gemini.requests.post
    sent = []

    def gemini_ok(url, headers=None, json=None, timeout=None):
        sent.append(json)
        return Reply(200, {"candidates": [{"content": {"parts": [{"text": '{"ok": true, "summary": "요약", "rules": ["규칙 하나"]}'}]}}]})

    def gemini_bad(url, headers=None, json=None, timeout=None):
        return Reply(400, text='{"error": {"message": "API key not valid. Please pass a valid API key."}}')

    try:
        os.environ.pop("GEMINI_API_KEY", None)
        check("키가 없으면 '키 없음'", _gemini.gemini_status()["state"] == "none")
        os.environ["GEMINI_API_KEY"] = "test-key-1"
        check("키를 넣은 직후에는 '연결 확인 전'", _gemini.gemini_status()["state"] == "unchecked")
        _gemini.requests.post = gemini_bad
        ok_flag, message = _gemini.GeminiAnalyzer(settings).check()
        check("틀린 키는 연결 실패로 표시", not ok_flag and _gemini.gemini_status()["state"] == "error" and "API 키" in message, message[:60])
        os.environ["GEMINI_API_KEY"] = "test-key-2"
        _gemini.requests.post = gemini_ok
        tested = app.post("/api/test_gemini", json={}, headers=host).get_json()
        check("연결되면 'Gemini API가 성공적으로 연결되었습니다' 안내와 연결됨 표시", tested["ok"] and "Gemini API가 성공적으로 연결되었습니다" in tested["message"]
              and tested["gemini"]["state"] == "ok" and "Connected" in tested["gemini"]["label"], tested["message"])
        check("상태 조회에도 연결됨으로 나옴", app.get("/api/state", headers=host).get_json()["state"]["gemini"]["connected"] is True)

        # 키 저장 -> 바로 연결 확인 (점검 중에는 .env 파일을 건드리지 않도록 저장 부분만 바꿔 끼운다)
        real_save = _config.save_env

        def memory_save(changes, path=None):
            for name, value in changes.items():
                if value:
                    os.environ[name] = value
                else:
                    os.environ.pop(name, None)

        _config.save_env = memory_save
        try:
            os.environ.pop("GEMINI_API_KEY", None)
            stored = app.post("/api/keys", json={"gemini_key": "test-key-3"}, headers=host).get_json()
            check("Gemini 키를 저장하면 바로 연결을 확인하고 AI 분석을 다시 요청", stored["ok"] and stored["gemini"]["ok"]
                  and "성공적으로 연결" in stored["gemini"]["message"] and stored["settings"]["gemini"]["state"] == "ok"
                  and oengine._flags["replan"] is True)
            oengine._flags["replan"] = False
            _gemini.requests.post = gemini_bad
            wrong_key = app.post("/api/keys", json={"gemini_key": "wrong-key"}, headers=host).get_json()
            check("틀린 키를 저장하면 연결 실패를 알려 줌", wrong_key["ok"] and wrong_key["gemini"]["ok"] is False
                  and wrong_key["settings"]["gemini"]["state"] == "error")
            _gemini.requests.post = gemini_ok
        finally:
            _config.save_env = real_save

        # Gemini가 형식이 어긋난 답을 줘도 안전하게 읽는다
        def gemini_odd(url, headers=None, json=None, timeout=None):
            answer = {"decisions": [{"symbol": "BTCUSDT", "action": "long", "confidence": 72, "p_long": 0.6, "p_short": 0.2},
                                    {"symbol": "ETHUSDT", "action": "short", "confidence": "NaN", "p_short": "abc"},
                                    "문자열 항목", {"symbol": "NOPEUSDT", "action": "long", "confidence": 0.9}]}
            return Reply(200, {"candidates": [{"content": {"parts": [{"text": _json.dumps(answer)}]}}]})

        _gemini.requests.post = gemini_odd
        odd = _gemini.GeminiAnalyzer(settings).decide({"indices": {}}, {}, None, [], [], {"pnl_pct": 0.0, "trades_done": 0, "loss_room": 25.0})
        by_symbol = {d["symbol"]: d for d in odd}
        check("AI 답의 숫자가 이상해도 안전하게 읽음 (72 -> 0.72, 확신도는 방향 확률 이하, 숫자가 아니면 0, 없는 코인은 버림)",
              set(by_symbol) == {"BTCUSDT", "ETHUSDT"} and by_symbol["BTCUSDT"]["confidence"] == 0.6
              and by_symbol["ETHUSDT"]["confidence"] == 0.0 and by_symbol["ETHUSDT"]["p_short"] is None)

        class Slow:
            calls = 0

        def gemini_limit(url, headers=None, json=None, timeout=None):
            Slow.calls += 1
            return Reply(429, text="quota")

        _gemini.requests.post = gemini_limit
        try:
            _gemini.GeminiAnalyzer(settings).client.generate_json("x")
            limited = ""
        except _gemini.GeminiError as exc:
            limited = str(exc)
        check("사용량 제한(429)이면 기다리며 반복하지 않고 바로 규칙 전략으로 넘어감", Slow.calls == 1 and "429" in limited
              and _gemini.gemini_status()["state"] == "ok" and _gemini.gemini_status()["warning"] is True)
        _gemini.requests.post = gemini_ok

        # 학습 자료
        os.environ.pop("GEMINI_API_KEY", None)
        kstorage = Storage(folder / "knowledge.db")
        ksettings = Settings(folder / "knowledge.json")
        kengine = TradeEngine(ksettings, kstorage, market=SyntheticMarket(seed=1, history_bars=400), analyzer=NoAI(), macro=NoMacro())
        kapp = create_app(kengine, ksettings, kstorage).test_client()
        added = kapp.post("/api/materials", json={"kind": "text", "title": "내 숏 전략", "text": "1시간봉 하락 추세에서 RSI 60 위로 반등하면 숏으로 진입한다."}, headers=host).get_json()
        check("전략 글 저장 (Gemini 키가 없으면 원문 그대로 지식 베이스에)", added["ok"] and added["materials"][0]["status"] == "pending"
              and "RSI 60 위로 반등하면 숏" in _knowledge.knowledge_text(kstorage))
        bad_url = kapp.post("/api/materials", json={"kind": "youtube", "url": "https://example.com/video"}, headers=host)
        short_text = kapp.post("/api/materials", json={"kind": "text", "text": "짧다"}, headers=host)
        big = kapp.post("/api/materials", json={"kind": "image", "mime": "image/png", "filename": "big.png",
                                                "image": base64.b64encode(b"0" * (7 * 1024 * 1024 + 1)).decode()}, headers=host)
        wrong = kapp.post("/api/materials", json={"kind": "image", "mime": "application/pdf", "image": "AAAA"}, headers=host)
        check("유튜브가 아닌 주소, 너무 짧은 글, 너무 큰 이미지, 이미지가 아닌 파일은 거부",
              [bad_url.status_code, short_text.status_code, big.status_code, wrong.status_code] == [400, 400, 400, 400])
        waiting = kapp.post("/api/materials", json={"kind": "youtube", "url": "https://www.youtube.com/watch?v=dQw4w9WgXcQ", "title": "강의"}, headers=host).get_json()
        check("유튜브 링크는 키가 없으면 분석 대기로 저장", waiting["ok"] and waiting["materials"][0]["status"] == "pending" and waiting["warning"])

        os.environ["GEMINI_API_KEY"] = "test-key-2"
        sent.clear()
        image = kapp.post("/api/materials", json={"kind": "image", "mime": "image/png", "filename": "chart.png", "title": "차트",
                                                  "image": base64.b64encode(b"\x89PNG fake image bytes").decode()}, headers=host).get_json()
        item = image["materials"][0]
        parts = sent[-1]["contents"][0]["parts"] if sent else []
        check("차트 이미지를 Gemini에 함께 보내 규칙으로 정리", image["ok"] and item["status"] == "ok" and item["rules"] == ["규칙 하나"]
              and any("inline_data" in part and part["inline_data"]["mime_type"] == "image/png" for part in parts))
        check("올린 이미지는 data/materials 폴더에 보관", (kstorage.folder / "materials" / f"{item['id']}.png").exists())
        retried = kapp.post("/api/materials/update", json={"id": waiting["materials"][0]["id"], "retry": True}, headers=host).get_json()
        video = next(m for m in retried["materials"] if m["kind"] == "youtube")
        video_parts = sent[-1]["contents"][0]["parts"]
        check("유튜브 링크를 Gemini에 영상으로 넘겨 분석", video["status"] == "ok"
              and any(part.get("file_data", {}).get("file_uri", "").startswith("https://www.youtube.com/") for part in video_parts))
        analyzer = _gemini.GeminiAnalyzer(ksettings)
        analyzer.knowledge = lambda: _knowledge.knowledge_text(kstorage)
        sent.clear()
        try:
            analyzer.daily_plan({"indices": {}}, {}, [], None)
        except Exception:
            pass
        prompt = sent[-1]["contents"][0]["parts"][-1]["text"] if sent else ""
        check("학습 자료의 규칙이 AI 프롬프트에 들어감", "규칙 하나" in prompt and "사용자가 직접 올린 학습 자료" in prompt and "RSI 60 위로 반등하면 숏" in prompt)
        off = kapp.post("/api/materials/update", json={"id": item["id"], "enabled": False}, headers=host).get_json()
        check("자료를 끄면 AI 판단에서 빠짐", off["in_use"] == len(off["materials"]) - 1)
        gone = kapp.post("/api/materials/update", json={"id": item["id"], "delete": True}, headers=host).get_json()
        check("자료 삭제 (이미지 파일도 함께)", all(m["id"] != item["id"] for m in gone["materials"])
              and not (kstorage.folder / "materials" / f"{item['id']}.png").exists())
    finally:
        _gemini.requests.post = real_post
        if saved_key is None:
            os.environ.pop("GEMINI_API_KEY", None)
        else:
            os.environ["GEMINI_API_KEY"] = saved_key

    print("\n13. 원하는 코인 직접 추가")
    csettings = Settings(folder / "coins.json")
    csettings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "entry_mode": "single", "max_hold_minutes": 60})
    cstorage2 = Storage(folder / "coins.db")
    cstart = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
    cclock = Clock(cstart)
    cmarket = SyntheticMarket(seed=8, history_bars=3000, start_ms=int(cstart.timestamp() * 1000), trend=0.00025)
    cmarket.add_symbol("ADAUSDT", 0.27, qty_step=1.0, tick=0.0001)       # Bybit에 ADA가 있는 상황을 흉내 낸다
    cengine = TradeEngine(csettings, cstorage2, market=cmarket, clock=cclock, analyzer=NoAI(), macro=NoMacro())
    capp = create_app(cengine, csettings, cstorage2).test_client()
    try:
        unknown = capp.post("/api/coins", json={"add": "nocoin"}, headers=host)
        weird = capp.post("/api/coins", json={"add": "!!"}, headers=host)
        check("Bybit에 없는 코인이나 잘못된 이름은 추가되지 않음", unknown.status_code == 400 and weird.status_code == 400 and "NOCOINUSDT" not in SYMBOLS,
              unknown.get_json()["error"][:40])
        added = capp.post("/api/coins", json={"add": "ada", "name": "에이다"}, headers=host).get_json()
        check("코인 추가 (약자만 넣어도 됨)", added["ok"] and SYMBOLS.get("ADAUSDT", {}).get("name") == "에이다"
              and "ADAUSDT" in csettings.get("trade_symbols") and any(c["symbol"] == "ADAUSDT" and c["custom"] for c in added["symbols"]),
              added.get("message", ""))
        again = capp.post("/api/coins", json={"add": "ADAUSDT"}, headers=host)
        base = capp.post("/api/coins", json={"remove": "BTCUSDT"}, headers=host)
        check("이미 있는 코인은 다시 추가되지 않고, 기본 코인은 지울 수 없음", again.status_code == 400 and base.status_code == 400)
        capp.post("/api/settings", json={"trade_symbols": ["ADAUSDT"]}, headers=host)
        for _ in range(288):
            cengine.step(); cmarket.advance(1); cclock.now += timedelta(minutes=5)
        traded = cstorage2.closed_trades("paper") + cstorage2.open_trades("paper")
        check("고른 코인(에이다)만 매수", bool(traded) and {t["symbol"] for t in traded} == {"ADAUSDT"}, f"{len(traded)}건 모두 ADA")
        page = capp.get("/", headers=host)
        chart = capp.get("/api/price?symbol=ADAUSDT&tf=15", headers=host).get_json()
        check("추가한 코인이 화면과 차트에 나옴", page.status_code == 200 and "ADAUSDT" in page.get_data(as_text=True) and chart["ok"] and chart["candles"])
        reloaded = Settings(folder / "coins.json")
        check("프로그램을 다시 켜도 추가한 코인이 유지됨", reloaded.get("custom_symbols") == [{"symbol": "ADAUSDT", "name": "에이다"}]
              and reloaded.get("trade_symbols") == ["ADAUSDT"])
        if not cstorage2.open_trades("paper"):          # 보유 포지션이 생길 때까지 조금 더 돌린다
            for _ in range(288):
                cengine.step(); cmarket.advance(1); cclock.now += timedelta(minutes=5)
                if cstorage2.open_trades("paper"):
                    break
        if cstorage2.open_trades("paper"):
            blocked = capp.post("/api/coins", json={"remove": "ADAUSDT"}, headers=host)
            check("보유 중인 코인은 목록에서 지울 수 없음", blocked.status_code == 400 and "ADAUSDT" in SYMBOLS)
            cengine.request_close_all(); cengine.step()
        only = capp.post("/api/coins", json={"remove": "ADAUSDT"}, headers=host)
        check("하나만 골라 둔 코인은 바로 지울 수 없음 (다른 코인이 말없이 켜지지 않게)", only.status_code == 400 and "ADAUSDT" in SYMBOLS
              and csettings.get("trade_symbols") == ["ADAUSDT"], only.get_json()["error"][:34])
        capp.post("/api/settings", json={"trade_symbols": ["BTCUSDT", "ADAUSDT"]}, headers=host)
        removed = capp.post("/api/coins", json={"remove": "ADAUSDT"}, headers=host).get_json()
        check("코인을 지우면 거래할 코인에서도 빠짐", removed["ok"] and "ADAUSDT" not in SYMBOLS and csettings.get("trade_symbols") == ["BTCUSDT"])
        hostile = Settings(folder / "hostile.json")
        (folder / "hostile.json").write_text('{"custom_symbols": [{"symbol": "ABCUSDT", "name": "<img src=x onerror=alert(1)>"}], '
                                             '"leverage_min": 40, "leverage_max": 10, "trade_symbols": 5}', encoding="utf-8")
        hostile.load()
        check("손으로 고친 설정 파일의 이상한 값은 걸러 냄 (기호가 든 코인 이름, 뒤바뀐 레버리지 범위)",
              "<" not in hostile.get("custom_symbols")[0]["name"] and hostile.get("leverage_min") <= hostile.get("leverage_max")
              and hostile.get("trade_symbols") == ["BTCUSDT", "ETHUSDT", "SOLUSDT", "XRPUSDT", "DOGEUSDT"])
        _config.apply_custom_symbols([])
        report = build_report(cstorage2, "paper", "1m", "Asia/Seoul", now=cclock.now)
        check("지운 코인의 지난 거래 기록은 그대로 남음", report["summary"]["trades"] == len(cstorage2.closed_trades("paper")) > 0
              and any(c["symbol"] == "ADAUSDT" and c["trades"] > 0 for c in report["by_coin"]) and report["trades"][0]["ticker"] == "ADA")
        cengine.step()
        check("코인을 지운 뒤에도 엔진과 화면이 정상", capp.get("/api/state", headers=host).status_code == 200
              and capp.get("/api/board", headers=host).status_code == 200)
    finally:
        _config.apply_custom_symbols([])


def section_v10(folder: Path, host: dict) -> None:
    """버전 10: Gemini 서버가 혼잡할 때(HTTP 503)의 대응."""
    import os

    import gemini_analyzer as _gemini
    from strategy import rule_plan

    print("\n14. Gemini 혼잡(503) 대응: 예비 모델, 계획 자동 전환, 모델 선택")

    class Reply:
        def __init__(self, code=200, data=None, text=""):
            self.status_code, self._data, self.text = code, data, text

        def json(self):
            if self._data is None:
                raise ValueError("no json")
            return self._data

    def model_row(name, methods=("generateContent",)):
        return {"name": "models/" + name, "displayName": name, "supportedGenerationMethods": list(methods)}

    listing = {"models": [model_row("gemini-9.9-flash-preview"), model_row("gemini-9.0-pro"), model_row("gemini-9.0-flash-lite"),
                          model_row("gemini-9.0-flash"), model_row("gemini-9.0-flash-image"),
                          model_row("text-embedding-004", ("embedContent",))]}
    busy, missing, calls = set(), set(), []

    def fake_get(url, headers=None, params=None, timeout=None):
        return Reply(200, listing)

    def fake_post(url, headers=None, json=None, timeout=None):
        model = url.split("/models/")[1].split(":")[0]
        calls.append(model)
        if model in missing:
            return Reply(404, text="not found")
        if model in busy:
            return Reply(503, text="The model is overloaded. Please try again later.")
        return Reply(200, {"candidates": [{"content": {"parts": [{"text": '{"ok": true}'}]}}]})

    saved_key = os.environ.get("GEMINI_API_KEY")
    real_post, real_get, real_sleep = _gemini.requests.post, _gemini.requests.get, _gemini.time.sleep
    _gemini.requests.post, _gemini.requests.get = fake_post, fake_get
    _gemini.time.sleep = lambda seconds: None          # 점검에서는 재시도 사이에 기다리지 않는다
    os.environ["GEMINI_API_KEY"] = "test-key-v10"
    _gemini._MODELS_CACHE.update(key="", at=0.0, items=[])
    try:
        primary = "gemini-9.9-flash-preview"
        client = _gemini.GeminiClient(primary)
        names = [m["name"] for m in client.list_models()]
        check("모델 목록: 글 생성 모델만, 정식 Flash 모델을 맨 앞에 (이미지, 임베딩 모델 제외)",
              names == ["gemini-9.0-flash", "gemini-9.0-flash-lite", "gemini-9.0-pro", "gemini-9.9-flash-preview"], ", ".join(names))

        busy.add(primary)
        result = client.generate_json("x")
        status = _gemini.gemini_status()
        check("기본 모델이 혼잡(503)하면 한 번 더 시도한 뒤 예비 모델로 넘어가 답을 받음",
              result == {"ok": True} and calls == [primary, primary, "gemini-9.0-flash"], " > ".join(calls))
        check("예비 모델로 답을 받으면 연결됨으로 두고, 예비 모델을 썼다는 안내를 남김",
              status["state"] == "ok" and not status.get("warning") and "예비 모델" in status["note"] and status["model"] == "gemini-9.0-flash")
        calls.clear()
        client.generate_json("x")
        check("혼잡이 이어지는 동안에는 응답했던 예비 모델을 먼저 부름 (기다리는 시간 없음)", calls == ["gemini-9.0-flash"], " > ".join(calls))

        busy.clear()
        client._detour = (client._detour[0], 0.0)          # 우회 시간이 지난 것으로 본다
        calls.clear()
        client.generate_json("x")
        check("시간이 지나면 기본 모델로 돌아가고 안내도 지움", calls == [primary] and _gemini.gemini_status()["note"] == "" and client._detour is None)

        busy.update(names)
        calls.clear()
        try:
            client.generate_json("x")
            message, code = "", None
        except _gemini.GeminiError as exc:
            message, code = str(exc), exc.status
        check("모든 모델이 혼잡하면 오래 붙잡지 않고(기본 2번 + 예비 2개) 규칙 전략으로 넘어감",
              code == 503 and len(calls) == 4 and "API 키나 설정 문제가 아닙니다" in message, f"호출 {len(calls)}번")
        check("혼잡으로 실패해도 키는 연결됨으로 두고 최근 호출 실패만 표시",
              _gemini.gemini_status()["state"] == "ok" and _gemini.gemini_status()["warning"] is True)

        busy.clear()
        busy.add(primary)
        picked = _gemini.GeminiClient(primary, backup="gemini-9.0-pro")
        calls.clear()
        picked.generate_json("x")
        check("예비 모델을 직접 정해 두면 그 모델부터 부름", calls[-1] == "gemini-9.0-pro", " > ".join(calls))

        busy.clear()
        missing.add("gemini-old")
        moved = _gemini.GeminiClient("gemini-old")
        moved.generate_json("x")
        replaced = moved.last_model
        moved.model = "gemini-9.0-pro"                     # 설정 화면에서 모델을 바꾼 경우
        moved.generate_json("x")
        check("없어진 모델은 쓸 수 있는 모델로 대신하고, 설정에서 모델을 바꾸면 바꾼 모델을 씀",
              replaced == "gemini-9.0-flash" and moved.last_model == "gemini-9.0-pro", f"{replaced} -> {moved.last_model}")

        no_key, failed = rule_plan({}, {"available": False}), rule_plan({}, {"available": False}, reason="failed")
        check("규칙 계획의 안내 문구가 실제 원인과 맞음 (키가 없을 때 / 호출이 실패했을 때)",
              "키가 없어" in no_key["cautions"][0] and "키가 없어" not in failed["cautions"][0]
              and "호출이 실패" in failed["cautions"][0] and failed["fallback"] is True and no_key["fallback"] is False)

        # 엔진: 계획 호출이 실패하면 규칙 계획으로 시작하고, Gemini가 응답하면 AI 계획으로 바꾼다
        class FlakyAI(FakeAI):
            def __init__(self):
                super().__init__(action="hold", confidence=0.2, p_long=0.2, p_short=0.2)
                self.down, self.plans = True, 0

            def daily_plan(self, macro, snapshots, lessons, last_review):
                self.plans += 1
                if self.down:
                    raise _gemini.GeminiError("Gemini 호출 실패 (Gemini 서버 오류 (HTTP 503))", 503)
                return super().daily_plan(macro, snapshots, lessons, last_review)

            def decide(self, *args, **kwargs):
                if self.down:
                    raise _gemini.GeminiError("Gemini 호출 실패 (Gemini 서버 오류 (HTTP 503))", 503)
                return super().decide(*args, **kwargs)

        flaky = FlakyAI()
        fsettings = Settings(folder / "flaky.json")
        fsettings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "ai_interval_minutes": 15})
        fstorage = Storage(folder / "flaky.db")
        start = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
        fclock = Clock(start)
        fmarket = SyntheticMarket(seed=3, history_bars=600, start_ms=int(start.timestamp() * 1000))
        fengine = TradeEngine(fsettings, fstorage, market=fmarket, clock=fclock, analyzer=flaky, macro=NoMacro())
        fapp = create_app(fengine, fsettings, fstorage).test_client()
        from utils import local_day

        def plan_now() -> dict:
            return fstorage.get_day("paper", local_day(fclock.now, "Asia/Seoul"))["plan"]

        fengine.step()
        first = fapp.get("/api/analysis", headers=host).get_json()
        check("계획 호출이 실패하면 규칙 계획으로 시작하고, 화면에 실패 사유와 시각을 알림 (같은 순환에서 두 번 부르지 않음)",
              plan_now()["source"] == "rule" and "호출이 실패" in plan_now()["cautions"][0]
              and "판단 실패" in first["ai_error"] and first["ai_error_at"] != "" and flaky.plans == 1, f"계획 호출 {flaky.plans}번")
        fmarket.advance(3)
        fclock.now += timedelta(minutes=15)
        fengine.step()
        check("Gemini가 계속 응답하지 않으면 규칙 계획을 유지하고, 실패하는 호출을 더 늘리지 않음", plan_now()["source"] == "rule" and flaky.plans == 1)
        flaky.down = False
        fmarket.advance(3)
        fclock.now += timedelta(minutes=15)
        fengine.step()
        later = fapp.get("/api/analysis", headers=host).get_json()
        check("Gemini가 다시 응답하면 오늘의 계획을 자동으로 AI 계획으로 바꿈", plan_now()["source"] == "gemini" and later["ai_error"] == ""
              and later["ai_error_at"] == "", plan_now()["source"])
        plans_before = flaky.plans
        fmarket.advance(3)
        fclock.now += timedelta(minutes=15)
        fengine.step()
        check("AI 계획으로 바뀐 뒤에는 계획을 다시 부르지 않음", flaky.plans == plans_before)

        # 화면: 모델 목록과 모델 저장
        msettings = Settings(folder / "models.json")
        mstorage = Storage(folder / "models.db")
        mengine = TradeEngine(msettings, mstorage, market=SyntheticMarket(seed=2, history_bars=400), analyzer=NoAI(), macro=NoMacro())
        mapp = create_app(mengine, msettings, mstorage).test_client()
        listed = mapp.get("/api/gemini_models", headers=host).get_json()
        check("API 화면용 모델 목록 조회", listed["ok"] and listed["models"][0]["name"] == "gemini-9.0-flash" and listed["backup"] == "auto"
              and listed["models"][0]["stable"] is True and listed["models"][-1]["stable"] is False)
        stored = mapp.post("/api/settings", json={"gemini_model": "gemini-9.0-flash", "gemini_backup_model": ""}, headers=host).get_json()
        wrong = mapp.post("/api/settings", json={"gemini_model": "모델 아님!"}, headers=host)
        check("모델과 예비 모델 저장 (예비를 비우면 자동, 이상한 이름은 거부)", stored["ok"] and stored["settings"]["gemini_model"] == "gemini-9.0-flash"
              and stored["settings"]["gemini_backup_model"] == "auto" and wrong.status_code == 400)
        os.environ.pop("GEMINI_API_KEY", None)
        empty = mapp.get("/api/gemini_models", headers=host).get_json()
        check("키가 없으면 모델 목록 대신 안내를 돌려줌", empty["ok"] and empty["models"] == [] and "키" in empty["message"])

        print("\n15. 기본 코인에 솔라나 추가")
        import config as _config
        check("기본 코인은 비트코인, 이더리움, 솔라나, 리플, 도지코인 순서",
              list(_config.BASE_SYMBOLS) == ["BTCUSDT", "ETHUSDT", "SOLUSDT", "XRPUSDT", "DOGEUSDT"]
              and Settings(folder / "fresh.json").get("trade_symbols") == list(_config.BASE_SYMBOLS))
        (folder / "old_coins.json").write_text('{"settings_version": 5, "trade_symbols": ["BTCUSDT", "DOGEUSDT"], '
                                               '"custom_symbols": [{"symbol": "SOLUSDT", "name": "솔라나"}]}', encoding="utf-8")
        old = Settings(folder / "old_coins.json")
        check("예전 설정을 불러오면 거래할 코인에 솔라나를 한 번 넣고, 직접 추가했던 솔라나는 기본 코인으로 합침",
              old.get("trade_symbols") == ["BTCUSDT", "SOLUSDT", "DOGEUSDT"] and old.get("custom_symbols") == [], str(old.get("trade_symbols")))
        (folder / "new_coins.json").write_text('{"settings_version": 6, "trade_symbols": ["BTCUSDT"]}', encoding="utf-8")
        check("새 버전에서 솔라나를 꺼 두면 다시 켜지 않음", Settings(folder / "new_coins.json").get("trade_symbols") == ["BTCUSDT"])
        _config.apply_custom_symbols([])
        sol_page = mapp.get("/api/price?symbol=SOLUSDT&tf=15", headers=host).get_json()
        removal = mapp.post("/api/coins", json={"remove": "SOLUSDT"}, headers=host)
        check("솔라나 차트가 나오고, 기본 코인이라 목록에서 지울 수 없음", sol_page["ok"] and bool(sol_page["candles"]) and removal.status_code == 400)

        print("\n16. 하루 목표와 청산을 투자금 기준으로 정하기")
        from risk_manager import daily_limit_text, daily_limits
        isettings = Settings(folder / "invest.json")
        saved = mapp.post("/api/settings", json={"invest_pct": 10, "daily_target_unit": "invest", "daily_target_invest_pct": 8,
                                                 "daily_loss_unit": "invest", "daily_loss_invest_pct": 9}, headers=host).get_json()
        got = daily_limits(msettings.all(), 1000.0)
        check("투자금 비율은 '오늘 원금 x 투자율'에 대한 비율 (원금 1,000, 투자율 10%, 목표 8%면 +8 USDT = 원금의 +0.8%)",
              saved["ok"] and abs(got[0] - 0.8) < 1e-9 and abs(got[1] - 0.9) < 1e-9, f"목표 +{got[0]:g}%, 청산 -{got[1]:g}%")
        texts = daily_limit_text(msettings.all())
        shown = mapp.get("/api/state", headers=host).get_json()["state"]
        check("화면 문구에 원금 대비 값과 투자금 대비 값을 함께 표시", texts == ("+0.8% (투자금의 +8%)", "-0.9% (투자금의 -9%)")
              and shown["target_text"] == texts[0] and shown["target_unit"] == "invest", " / ".join(texts))
        isettings.update({"invest_pct": 100, "daily_loss_unit": "invest", "daily_loss_invest_pct": 500})
        check("투자금 기준 손실 한도도 원금 전체(-100%)를 넘지 않음", daily_limits(isettings.all(), 1000.0)[1] == 100.0)
        wrong_unit = mapp.post("/api/settings", json={"daily_target_unit": "margin"}, headers=host)
        check("알 수 없는 방식은 저장하지 않음", wrong_unit.status_code == 400)
        lsettings, lstorage, lengine, lclock = simulate(folder, "invest_lock", 2, {"invest_pct": 60, "leverage_mode": "manual", "manual_leverage": 10,
                                                                               "daily_loss_unit": "invest", "daily_loss_invest_pct": 3})
        locked = [d for d in lstorage.days_between("paper", "2026-01-01", "2027-01-01") if d["locked"]]
        check("투자금 기준 손실 한도에 닿으면 그날 잠기고 사유에 기준을 적음", bool(locked) and "투자금의 -3%" in (locked[0]["lock_reason"] or ""),
              (locked[0]["lock_reason"] if locked else "잠긴 날 없음"))

        print("\n17. 보유 포지션의 익절/손절 직접 수정 (투자금 기준)")
        from broker import LiveBroker
        from bybit_api import BybitError
        psettings = Settings(folder / "stops.json")
        psettings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "ai_interval_minutes": 15, "entry_mode": "single",
                          "leverage_mode": "manual", "manual_leverage": 5, "invest_pct": 50, "max_hold_minutes": 1440,
                          "daily_loss_limit_pct": 100, "paper_start_balance": 1000.0})
        pstorage = Storage(folder / "stops.db")
        pstart = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
        pclock = Clock(pstart)
        pmarket = SyntheticMarket(seed=5, history_bars=3000, start_ms=int(pstart.timestamp() * 1000))
        pengine = TradeEngine(psettings, pstorage, market=pmarket, clock=pclock, analyzer=NoAI(), macro=NoMacro())
        papp = create_app(pengine, psettings, pstorage).test_client()
        for _ in range(900):
            pengine.step()
            if pstorage.open_trades("paper"):
                break
            pmarket.advance(1)
            pclock.now += timedelta(minutes=5)
        opened = pstorage.open_trades("paper")
        if not opened:
            check("점검용 포지션 만들기", False, "포지션이 열리지 않았습니다")
        else:
            before = opened[0]
            tid, margin, long = before["id"], before["margin"], before["side"] == "long"

            def call(**body):
                return papp.post("/api/position_stops", json={"id": tid, **body}, headers=host)

            def position():
                return next(p for p in papp.get("/api/state", headers=host).get_json()["state"]["positions"] if p["id"] == tid)

            shown = position()
            check("자동으로 걸린 익절/손절이 투자금 대비 %로 표시됨 (기본은 자동)", shown["tpsl_mode"] == "auto" and shown["tpsl_edited"] is False
                  and shown["tp_margin_pct"] > 0 > shown["sl_margin_pct"], f"익절 {shown['tp_margin_pct']:+.1f}% / 손절 {shown['sl_margin_pct']:+.1f}%")
            wanted_tp = round(shown["tp_margin_pct"] + 10, 1)
            preview = call(tp={"unit": "pct", "value": wanted_tp}, preview=True).get_json()
            check("미리 보기: 익절만 바꾸면 손절은 그대로, 저장은 하지 않음", preview["ok"] and 0 <= preview["stops"]["tp_margin_pct"] - wanted_tp < 0.5
                  and preview["stops"]["sl_price"] == before["sl_price"] and preview["stops"]["applied"] is False
                  and pstorage.open_trades("paper")[0]["tp_price"] == before["tp_price"], f"익절 {preview['stops']['tp_margin_pct']:+.2f}%")
            saved_stops = call(tp={"unit": "pct", "value": wanted_tp}).get_json()
            after = pstorage.open_trades("paper")[0]
            pengine.step()                 # 실제로는 저장 직후 엔진이 깨어나 화면용 상태를 새로 만든다
            shown = position()
            check("익절을 투자금의 %로 직접 고쳐 저장 (가격으로 바꿔 포지션에 건다)", saved_stops["ok"] and saved_stops["stops"]["applied"]
                  and after["tp_price"] == saved_stops["stops"]["tp_price"] != before["tp_price"] and after["sl_price"] == before["sl_price"]
                  and abs(shown["tp_margin_pct"] - wanted_tp) < 0.5 and shown["tpsl_mode"] == "manual" and shown["tpsl_edited"] is True,
                  f"{saved_stops.get('message', '')}")
            loss_amount = round(margin * 0.05, 2)
            by_amount = call(sl={"unit": "amount", "value": loss_amount}).get_json()
            now_net = position()["unrealized"]
            if by_amount["ok"]:
                check("손절을 금액(USDT)으로 직접 고쳐 저장", abs(abs(by_amount["stops"]["sl_net"]) - loss_amount) < max(0.05, loss_amount * 0.02)
                      and pstorage.open_trades("paper")[0]["sl_price"] == by_amount["stops"]["sl_price"], f"손절 시 {by_amount['stops']['sl_net']:+.2f} USDT")
            else:
                check("손절을 금액(USDT)으로 직접 고쳐 저장 (지금 손실이 이미 그 값을 넘으면 거부)", now_net <= -loss_amount * 0.9, by_amount["error"][:40])
            too_wide = call(sl={"unit": "pct", "value": 80})
            not_number = call(tp={"unit": "pct", "value": "많이"})
            negative = call(tp={"unit": "pct", "value": -5})
            missing = papp.post("/api/position_stops", json={"id": 999999, "tp": {"unit": "pct", "value": 10}}, headers=host)
            check("잘못된 값은 저장하지 않고 이유를 알려 줌 (손절 60% 초과, 숫자 아님, 0 이하, 없는 포지션)",
                  [too_wide.status_code, not_number.status_code, negative.status_code, missing.status_code] == [400, 400, 400, 400]
                  and "60%" in too_wide.get_json()["error"], too_wide.get_json()["error"][:36])
            pnl_now = position()["unrealized_pct"]
            passed_tp = call(tp={"unit": "pct", "value": max(0.01, pnl_now / 2)}) if pnl_now > 0.5 else None
            passed_sl = call(sl={"unit": "pct", "value": abs(pnl_now) / 2}) if pnl_now < -0.5 else None
            hit = passed_tp or passed_sl
            check("이미 지나친 값(저장하면 바로 청산될 값)은 거부", hit is None or hit.status_code == 400,
                  hit.get_json()["error"][:30] if hit is not None else "지금 손익이 0에 가까워 건너뜀")
            # 바꾼 익절 가격에 닿으면 그 가격으로 익절된다
            target = pstorage.open_trades("paper")[0]["tp_price"]
            pmarket.bars[before["symbol"]][-1]["c"] = target * (1.002 if long else 0.998)
            pengine.step()
            closed = [t for t in pstorage.closed_trades("paper") if t["id"] == tid]
            check("직접 고친 익절 가격에 닿으면 익절로 청산", bool(closed) and closed[0]["exit_reason"] == "익절" and closed[0]["net_pnl"] > 0,
                  f"순손익 {closed[0]['net_pnl']:+.2f} USDT (투자금의 {closed[0]['net_pnl'] / margin * 100:+.1f}%)" if closed else "청산되지 않음")

        class RefusingClient:
            env = "demo"

            def set_trading_stop(self, symbol, stop_loss, take_profit):
                raise BybitError("Bybit 오류 10001: invalid price", 10001)

        refusing = LiveBroker(RefusingClient(), None, psettings)
        quiet = refusing.set_stops({"symbol": "BTCUSDT"}, 1.0, 2.0, {"tick_size": 0.1}) is None
        try:
            refusing.set_stops({"symbol": "BTCUSDT"}, 1.0, 2.0, {"tick_size": 0.1}, strict=True)
            raised = False
        except BybitError:
            raised = True
        check("실거래에서 거래소가 변경을 받지 않으면, 직접 수정할 때는 화면에 알림 (자동 조정 때는 기록만)", quiet and raised)

        print("\n18. 남은 금액과 직접 주문")
        osettings = Settings(folder / "order.json")
        osettings.update({"mode": "paper", "trading_enabled": False, "timezone": "Asia/Seoul", "paper_start_balance": 1000.0,
                          "max_hold_minutes": 1440, "daily_loss_limit_pct": 100})
        ostorage = Storage(folder / "order.db")
        ostart = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
        oclock = Clock(ostart)
        omarket = SyntheticMarket(seed=9, history_bars=3000, start_ms=int(ostart.timestamp() * 1000))
        oengine2 = TradeEngine(osettings, ostorage, market=omarket, clock=oclock, analyzer=NoAI(), macro=NoMacro())
        oapp = create_app(oengine2, osettings, ostorage).test_client()
        oengine2.step()

        def order(**body):
            return oapp.post("/api/manual_order", json=body, headers=host)

        def ostate():
            oengine2.step()
            return oapp.get("/api/state", headers=host).get_json()["state"]

        first_state = ostate()
        check("포지션이 없으면 남은 금액 = 총자산 (모의투자도 같은 자리에 표시)", abs(first_state["available"] - first_state["equity"]) < 1e-9
              and first_state["invested"] == 0 and abs(first_state["available"] - 1000.0) < 1e-6, f"남은 금액 {first_state['available']:.2f} USDT")
        seen = order(symbol="BTCUSDT", side="long", margin=250, leverage=5, preview=True).get_json()
        check("주문 미리 보기: 수량, 거래 금액, 자동 익절/손절, 주문 뒤 남는 금액을 계산만 함", seen["ok"] and 250 - seen["order"]["price"] * seen["order"]["qty_step"] / 5 - 1e-6 <= seen["order"]["margin"] <= 250 + 1e-6
              and abs(seen["order"]["notional"] - seen["order"]["margin"] * 5) < 1e-6 and seen["order"]["tp_auto"] and seen["order"]["sl_auto"]
              and seen["order"]["tp_margin_pct"] > 0 > seen["order"]["sl_margin_pct"] and not ostorage.open_trades("paper")
              and abs(seen["order"]["left_after"] - (1000 - seen["order"]["margin"] - seen["order"]["fee_in"])) < 1e-6,
              f"투자금 {seen['order']['margin']:.2f}, 남는 금액 {seen['order']['left_after']:.2f}")
        placed = order(symbol="BTCUSDT", side="long", margin=250, leverage=5).get_json()
        held = ostorage.open_trades("paper")
        after_state = ostate()
        check("직접 주문이 시장가로 체결되고 보유 포지션에 '직접 주문'으로 표시 (자동 매매가 꺼져 있어도)", placed["ok"] and placed["order"]["applied"]
              and len(held) == 1 and held[0]["source"] == "manual" and held[0]["leverage"] == 5 and after_state["positions"][0]["manual"] is True,
              placed.get("message", ""))
        check("남은 금액 = 총자산 - 투자 중인 증거금", abs(after_state["available"] - (after_state["equity"] - held[0]["margin"])) < 1e-6
              and abs(after_state["invested"] - held[0]["margin"]) < 1e-9, f"투자 중 {after_state['invested']:.2f}, 남은 금액 {after_state['available']:.2f}")
        noted = ostorage.recent_analysis("paper", 5)
        check("직접 주문도 분석 기록에 남음", bool(noted) and noted[0]["source"] == "manual" and noted[0]["executed"] == 1 and "투자금" in noted[0]["note"])
        again = order(symbol="BTCUSDT", side="short", margin=50, leverage=5)
        too_much = order(symbol="ETHUSDT", side="long", margin=5000, leverage=5)
        too_small = order(symbol="ETHUSDT", side="long", margin=0.01, leverage=1)
        bad_lev = order(symbol="ETHUSDT", side="long", margin=50, leverage=500)
        bad_coin = order(symbol="NOPEUSDT", side="long", margin=50, leverage=5)
        tight = order(symbol="ETHUSDT", side="long", margin=50, leverage=50, sl={"unit": "pct", "value": 1})
        check("받을 수 없는 주문은 이유를 알려 줌 (같은 코인 보유, 남은 금액 초과, 최소 주문 미달, 배율 범위, 없는 코인, 수수료보다 좁은 손절)",
              [r.status_code for r in (again, too_much, too_small, bad_lev, bad_coin, tight)] == [400] * 6
              and "이미 보유" in again.get_json()["error"] and "남은 금액" in too_much.get_json()["error"]
              and "수수료" in tight.get_json()["error"], too_much.get_json()["error"][:44])
        custom = order(symbol="ETHUSDT", side="short", margin=100, leverage=5, tp={"unit": "pct", "value": 20}, sl={"unit": "pct", "value": 10}).get_json()
        eth = next(p for p in ostate()["positions"] if p["symbol"] == "ETHUSDT")
        check("익절/손절을 직접 넣어 주문 (투자금의 +20% / -10%)", custom["ok"] and abs(eth["tp_margin_pct"] - 20) < 1.0 and abs(eth["sl_margin_pct"] + 10) < 1.0
              and eth["tpsl_edited"] is True and eth["side"] == "short", f"익절 {eth['tp_margin_pct']:+.1f}% / 손절 {eth['sl_margin_pct']:+.1f}%")
        osettings.update({"mode": "live"})
        unconfirmed = order(symbol="XRPUSDT", side="long", margin=50, leverage=5)
        osettings.update({"mode": "paper"})
        check("실전 거래 모드의 주문은 확인을 거쳐야 함", unconfirmed.status_code == 400 and "확인" in unconfirmed.get_json()["error"])
        btc = ostorage.open_trades("paper")[0]
        omarket.bars["BTCUSDT"][-1]["c"] = btc["tp_price"] * 1.002
        oengine2.step()
        done = [t for t in ostorage.closed_trades("paper") if t["id"] == btc["id"]]
        check("직접 넣은 포지션도 익절/손절 가격에 닿으면 자동으로 청산 (자동 매매가 꺼져 있어도 감시)", bool(done) and done[0]["exit_reason"] == "익절")
        ostorage.update_day("paper", local_day(oclock.now, "Asia/Seoul"), locked=1, lock_reason="점검")
        locked_out = order(symbol="XRPUSDT", side="long", margin=50, leverage=5)
        check("하루 청산 한도로 잠긴 날에는 직접 주문도 받지 않음", locked_out.status_code == 400 and "멈춰" in locked_out.get_json()["error"])

        print("\n19. 기록의 달력: 날짜별 수익, 손실, 분석")
        csettings2, cstorage3, cengine2, cclock2 = simulate(folder, "calendar", 2, {"invest_pct": 30, "leverage_mode": "manual", "manual_leverage": 5})
        calapp = create_app(cengine2, csettings2, cstorage3).test_client()
        cal = calapp.get("/api/calendar?month=2026-09", headers=host).get_json()
        with_trades = [d for d in cal["days"] if d["trades"]]
        check("달력: 그 달의 날짜별 손익과 기록 유무 (9월은 30일, 1일은 화요일)", cal["ok"] and cal["days_in_month"] == 30 and cal["first_weekday"] == 2
              and bool(with_trades) and all(abs(d["profit"] + d["loss"] - d["net"]) < 1e-9 for d in cal["days"]) and cal["next"] == "2026-10",
              f"거래가 있는 날 {len(with_trades)}일")
        pick = max(cal["days"], key=lambda d: d["analyses"])["day"]
        one = calapp.get(f"/api/day?date={pick}", headers=host).get_json()
        rows_db = cstorage3.trades_on_day("paper", pick)
        check("날짜를 고르면 그날의 수익 합계, 손실 합계, 순손익, 거래를 보여 줌", one["ok"] and one["summary"]["trades"] == len(rows_db)
              and abs(one["summary"]["net"] - sum(t["net_pnl"] for t in rows_db)) < 1e-6 and one["summary"]["profit"] >= 0 >= one["summary"]["loss"]
              and abs(one["summary"]["profit"] + one["summary"]["loss"] - one["summary"]["net"]) < 1e-9 and len(one["trades"]) == len(rows_db)
              and one["plan"] is not None, f"{pick}: 수익 {one['summary']['profit']:+.2f}, 손실 {one['summary']['loss']:+.2f}, 거래 {len(rows_db)}회")
        total = one["analysis"]["total"]
        check("그날의 분석은 처음에 15개", total > 35 and len(one["analysis"]["rows"]) == 15 and one["analysis"]["has_more"] is True, f"그날 분석 {total}건")
        more = calapp.get(f"/api/day?date={pick}&offset=15&limit=20&only=analysis", headers=host).get_json()
        third = calapp.get(f"/api/day?date={pick}&offset=35&limit=20&only=analysis", headers=host).get_json()
        key = lambda r: (r["time"], r["symbol"], r["action"], r["reason"])
        loaded = one["analysis"]["rows"] + more["analysis"]["rows"] + third["analysis"]["rows"]
        check("'더 보기'를 누를 때마다 20개씩 이어서 나옴 (겹치거나 빠지지 않음)", len(more["analysis"]["rows"]) == 20 and "summary" not in more
              and len(third["analysis"]["rows"]) == min(20, total - 35) and len(loaded) == min(total, 55)
              and [key(r) for r in loaded] == [key({"time": fmt, "symbol": r["symbol"], "action": r["action"], "reason": r["reason"]})
                                              for r, fmt in ((r, __import__("utils").fmt_local(r["ts"], "Asia/Seoul", "%H:%M"))
                                                             for r in cstorage3.analysis_on_day("paper", pick, 55, 0))])
        last_page = calapp.get(f"/api/day?date={pick}&offset={max(0, total - 5)}&limit=20&only=analysis", headers=host).get_json()
        check("마지막까지 가면 더 보기가 사라짐", len(last_page["analysis"]["rows"]) == min(5, total) and last_page["analysis"]["has_more"] is False)
        empty_day = calapp.get("/api/day?date=2026-08-15", headers=host).get_json()
        future = calapp.get("/api/day?date=2999-01-01", headers=host)
        odd = calapp.get("/api/day?date=어제", headers=host)
        check("기록이 없는 날은 빈 화면, 미래나 이상한 날짜는 거부", empty_day["ok"] and empty_day["has_record"] is False and empty_day["summary"]["trades"] == 0
              and future.status_code == 400 and odd.status_code == 400)
        cstorage3.prune_analysis("paper")
        check("지난 날짜의 분석을 볼 수 있도록 판단 기록을 넉넉히 보관", cstorage3.count_analysis_on_day("paper", pick) == total)

        print("\n20. 동시 보유 포지션 5개, 다음 판단 시각")
        five = mapp.post("/api/settings", json={"max_open_positions": 5, "trade_symbols": list(_config.BASE_SYMBOLS)}, headers=host).get_json()
        six = mapp.post("/api/settings", json={"max_open_positions": 6}, headers=host)
        check("동시 보유 포지션을 5개까지 저장 (6개는 거부)", five["ok"] and five["settings"]["max_open_positions"] == 5 and six.status_code == 400
              and mengine.risk.slots() == 5)
        (folder / "old_slots.json").write_text('{"settings_version": 6, "max_open_positions": 4}', encoding="utf-8")
        (folder / "old_slots2.json").write_text('{"settings_version": 6, "max_open_positions": 2}', encoding="utf-8")
        (folder / "new_slots.json").write_text('{"settings_version": 7, "max_open_positions": 4}', encoding="utf-8")
        check("예전 최대값(4개)으로 쓰던 설정은 5개로 올리고, 그보다 작게 쓰던 값이나 새 버전에서 정한 4개는 그대로 둠",
              Settings(folder / "old_slots.json").get("max_open_positions") == 5 and Settings(folder / "old_slots2.json").get("max_open_positions") == 2
              and Settings(folder / "new_slots.json").get("max_open_positions") == 4)
        from risk_manager import RiskManager
        wide = RiskManager(msettings)
        msettings.update({"invest_pct": 100})
        check("포지션 5개면 투자 한도를 5로 나눈 금액이 포지션 하나의 투자금", abs(wide.planned_margin(1000.0) - 200.0) < 1e-9)
        nsettings = Settings(folder / "next.json")
        nsettings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "ai_interval_minutes": 30})
        nstorage = Storage(folder / "next.db")
        nclock = Clock(datetime.now(timezone.utc).replace(microsecond=0))
        nengine = TradeEngine(nsettings, nstorage, market=SyntheticMarket(seed=4, history_bars=600), clock=nclock, analyzer=NoAI(), macro=NoMacro())
        napp = create_app(nengine, nsettings, nstorage).test_client()
        nengine.step()
        ns = napp.get("/api/state", headers=host).get_json()["state"]
        check("운용 상태에 다음 판단 시각과 남은 시간, 판단 주기, 마지막 판단 시각을 알려 줌", ns["next_decision"] and ns["last_decision"]
              and ns["ai_interval_minutes"] == 30 and 29 * 60 <= ns["next_decision_in"] <= 30 * 60, f"다음 판단 {ns['next_decision']}, {ns['next_decision_in']}초 뒤")
    finally:
        _gemini.requests.post, _gemini.requests.get, _gemini.time.sleep = real_post, real_get, real_sleep
        _gemini._MODELS_CACHE.update(key="", at=0.0, items=[])
        if saved_key is None:
            os.environ.pop("GEMINI_API_KEY", None)
        else:
            os.environ["GEMINI_API_KEY"] = saved_key


class ExitAI(FakeAI):
    """매도 방식 점검용 가짜 AI. 새 진입에는 정해 둔 exit를, 보유 포지션에는 정해 둔 행동을 돌려준다."""

    def __init__(self, way=None, held_action="hold", entry="single", **kwargs):
        super().__init__(**kwargs)
        self.way, self.held_action, self.entry = way, held_action, entry

    def decide(self, macro, snapshots, plan, lessons, open_trades, state, history=None, rule_views=None, balance=None):
        rows = super().decide(macro, snapshots, plan, lessons, open_trades, state, history, rule_views, balance)
        held = {t["symbol"] for t in open_trades}
        for row in rows:
            row["entry"] = self.entry
            if row["symbol"] in held:
                row["action"] = self.held_action
            elif self.way:
                row["exit"] = self.way
        return rows


def section_v12(folder: Path, host: dict) -> None:
    """버전 12: 매도 방식 (AI 자동 / 일괄 매도 / 분할 매도)."""
    import json as _json
    import broker as _broker
    import gemini_analyzer as _gemini
    from broker import LiveBroker
    from strategy import rule_decide
    from market_data import all_snapshots

    print("\n21. 매도 방식: AI 자동 / 일괄 매도 / 분할 매도")
    fresh = Settings(folder / "exit_default.json")
    check("매도 방식의 기본값은 AI 자동, 분할 매도는 3회", fresh.get("exit_mode") == "auto" and fresh.get("exit_split_count") == 3)
    try:
        fresh.update({"exit_mode": "나중에"})
        refused = False
    except Exception:
        refused = True
    try:
        fresh.update({"exit_split_count": 9})
        refused_count = False
    except Exception:
        refused_count = True
    check("잘못된 매도 방식과 범위를 벗어난 횟수는 저장하지 않음 (2~5회)", refused and refused_count)

    def build(name, changes, analyzer=None, seed=5):
        settings = Settings(folder / f"{name}.json")
        settings.update({"mode": "paper", "trading_enabled": False, "timezone": "Asia/Seoul", "ai_interval_minutes": 15,
                         "entry_mode": "single", "leverage_mode": "manual", "manual_leverage": 5, "invest_pct": 50,
                         "max_hold_minutes": 1440, "daily_loss_limit_pct": 100, "paper_start_balance": 1000.0, **changes})
        storage = Storage(folder / f"{name}.db")
        start = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
        clock = Clock(start)
        market = SyntheticMarket(seed=seed, history_bars=3000, start_ms=int(start.timestamp() * 1000))
        engine = TradeEngine(settings, storage, market=market, clock=clock, analyzer=analyzer or NoAI(), macro=NoMacro())
        engine.step()
        return settings, storage, engine, clock, market

    def set_price(market, symbol, value):
        market.bars[symbol][-1]["c"] = value

    def tick(engine, clock, seconds=5):
        clock.now += timedelta(seconds=seconds)
        engine.step()

    # --- 분할 매도 고정: 가격이 익절가 쪽으로 구간을 지날 때마다 한 회차씩 판다
    ssettings, sstorage, sengine, sclock, smarket = build("exit_split", {"exit_mode": "split", "exit_split_count": 3})
    order = sengine.manual_order("BTCUSDT", "long", 200, 5, tp={"unit": "pct", "value": 20}, sl={"unit": "pct", "value": 10}, apply=True)
    sengine.step()
    held = sstorage.open_trades("paper")[0]
    tid, entry, tp, qty0, margin0 = held["id"], held["entry_price"], held["tp_price"], held["qty"], held["margin"]
    at = lambda share: entry + (tp - entry) * share
    shown = sengine.get_state()
    pos = shown["positions"][0]
    check("분할 매도 고정: 운용 상태와 보유 포지션에 매도 방식과 회차, 다음에 팔 가격이 나옴",
          shown["exit_mode"] == "split" and shown["exit_split_count"] == 3 and pos["exit_total"] == 3 and pos["exit_done"] == 0
          and abs(pos["exit_next"] - at(1 / 3)) < entry * 1e-6, f"다음 매도가 {pos['exit_next']:.2f} (진입 {entry:.2f}, 익절 {tp:.2f})")
    set_price(smarket, "BTCUSDT", at(0.2))
    tick(sengine, sclock)
    check("첫 구간(1/3)에 닿기 전에는 팔지 않음", sstorage.open_trades("paper")[0]["qty"] == qty0 and not sstorage.closed_trades("paper"))
    set_price(smarket, "BTCUSDT", at(0.4))
    before_equity = sengine._net_equity(sengine.broker_for("paper", smarket), smarket, "paper")[0]     # 팔기 직전, 같은 가격에서의 순자산
    tick(sengine, sclock)
    parts = sstorage.closed_trades("paper")
    left = sstorage.open_trades("paper")[0]
    after_equity = sengine.get_state()["equity"]
    check("첫 구간에 닿으면 1/3을 팔아 그 몫의 수익을 확정", len(parts) == 1 and parts[0]["exit_reason"].startswith("분할 매도 1/3")
          and parts[0]["net_pnl"] > 0 and abs(parts[0]["qty"] - qty0 / 3) <= 0.001 + 1e-9
          and abs(left["qty"] + parts[0]["qty"] - qty0) < 1e-9, f"{parts[0]['exit_reason'] if parts else '팔리지 않음'}")
    check("판 몫만큼 보유 수량과 투자금이 줄고, 손절가와 익절가는 그대로", abs(left["margin"] + parts[0]["margin"] - margin0) < 1e-6
          and left["tp_price"] == tp and left["sl_price"] == held["sl_price"] and left["id"] == tid
          and abs(left["fee"] - held["fee"] * left["qty"] / qty0) < 1e-9, f"투자금 {margin0:.2f} -> {left['margin']:.2f} USDT")
    check("나눠 팔아도 순자산은 그대로 (미실현 수익이 실현 수익으로 바뀔 뿐)", abs(after_equity - before_equity) < 0.02,
          f"{before_equity:.4f} -> {after_equity:.4f}")
    logged = [r for r in sstorage.recent_analysis("paper", 50) if r["action"] == "reduce"]
    check("분할 매도한 이유가 분석 기록에 남음", len(logged) == 1 and logged[0]["executed"] == 1 and "분할 매도 1/3회차" in logged[0]["note"],
          logged[0]["note"][:50] if logged else "기록 없음")
    tick(sengine, sclock)
    check("같은 구간에서는 다시 팔지 않음", len(sstorage.closed_trades("paper")) == 1)
    set_price(smarket, "BTCUSDT", at(0.7))
    tick(sengine, sclock)
    parts = sstorage.closed_trades("paper")
    left = sstorage.open_trades("paper")[0]
    check("둘째 구간(2/3)에 닿으면 한 회차를 더 팖", len(parts) == 2 and parts[1]["exit_reason"].startswith("분할 매도 2/3")
          and abs(left["qty"] - (qty0 - parts[0]["qty"] - parts[1]["qty"])) < 1e-9, f"남은 수량 {left['qty']:g} / 처음 {qty0:g}")
    pos = sengine.get_state()["positions"][0]
    check("마지막 회차는 익절가에서 정리하므로 다음 매도가는 따로 없음", pos["exit_done"] == 2 and pos["exit_next"] is None)
    set_price(smarket, "BTCUSDT", at(0.95))
    tick(sengine, sclock)
    check("마지막 회차는 익절가에 닿기 전에는 팔지 않음", len(sstorage.closed_trades("paper")) == 2)
    set_price(smarket, "BTCUSDT", tp * 1.001)
    tick(sengine, sclock)
    parts = sstorage.closed_trades("paper")
    whole = [t for t in parts if not is_part(t)]
    check("익절가에 닿으면 남은 수량이 익절로 정리됨", len(parts) == 3 and not sstorage.open_trades("paper")
          and len(whole) == 1 and whole[0]["id"] == tid and whole[0]["exit_reason"] == "익절")
    check("나눠 판 수량과 투자금의 합이 처음 포지션과 같음", abs(sum(t["qty"] for t in parts) - qty0) < 1e-9
          and abs(sum(t["margin"] for t in parts) - margin0) < 1e-6)
    balance = float(sstorage.kv_get("paper_balance"))
    check("세 번에 걸쳐 확정한 순손익의 합이 잔고 변화와 같음", abs(balance - 1000.0 - sum(t["net_pnl"] for t in parts)) < 1e-6,
          f"순손익 합 {sum(t['net_pnl'] for t in parts):+.4f} USDT")
    day = sengine.get_state()["day"]
    check("거래 횟수는 포지션 하나를 한 번으로 셈 (기록은 회차마다 남음)", sstorage.count_trades_on_day("paper", day) == 1
          and len(sstorage.trades_on_day("paper", day)) == 3 and len(sstorage.recent_entries("paper", "2000-01-01")) == 1)

    # --- 숏 포지션: 가격이 내려갈수록 익절가에 가까워진다
    hsettings, hstorage, hengine, hclock, hmarket = build("exit_short", {"exit_mode": "split", "exit_split_count": 2})
    hengine.manual_order("ETHUSDT", "short", 200, 5, tp={"unit": "pct", "value": 20}, sl={"unit": "pct", "value": 10}, apply=True)
    hengine.step()
    h = hstorage.open_trades("paper")[0]
    set_price(hmarket, "ETHUSDT", h["entry_price"] + (h["tp_price"] - h["entry_price"]) * 0.3)
    tick(hengine, hclock)
    none_yet = not hstorage.closed_trades("paper")
    set_price(hmarket, "ETHUSDT", h["entry_price"] + (h["tp_price"] - h["entry_price"]) * 0.6)
    tick(hengine, hclock)
    sold = hstorage.closed_trades("paper")
    check("숏 포지션도 익절가 쪽(아래)으로 절반을 가면 2회 중 1회를 팖", none_yet and len(sold) == 1 and sold[0]["side"] == "short"
          and sold[0]["net_pnl"] > 0 and sold[0]["exit_reason"].startswith("분할 매도 1/2"), sold[0]["exit_reason"] if sold else "팔리지 않음")
    set_price(hmarket, "ETHUSDT", h["sl_price"] * 1.001)
    tick(hengine, hclock)
    rest = [t for t in hstorage.closed_trades("paper") if not is_part(t)]
    check("일부를 판 뒤 손절가에 닿으면 남은 수량만 손절", len(rest) == 1 and rest[0]["exit_reason"] == "손절" and rest[0]["net_pnl"] < 0
          and abs(rest[0]["qty"] - (h["qty"] - sold[0]["qty"])) < 1e-9, f"남은 수량 {rest[0]['qty']:g} 손절 {rest[0]['net_pnl']:+.2f} USDT" if rest else "")

    # --- 일괄 매도 고정: 중간에 팔지 않는다
    osettings, ostorage, oengine, oclock, omarket = build("exit_single", {"exit_mode": "single"})
    oengine.manual_order("BTCUSDT", "long", 200, 5, tp={"unit": "pct", "value": 20}, sl={"unit": "pct", "value": 10}, apply=True)
    oengine.step()
    o = ostorage.open_trades("paper")[0]
    set_price(omarket, "BTCUSDT", o["entry_price"] + (o["tp_price"] - o["entry_price"]) * 0.8)
    tick(oengine, oclock)
    opos = oengine.get_state()["positions"][0]
    check("일괄 매도 고정: 익절가로 가는 중간에는 팔지 않음", not ostorage.closed_trades("paper") and opos["exit_total"] is None
          and ostorage.open_trades("paper")[0]["qty"] == o["qty"])
    blocked, why = oengine._reduce_by_decision("paper", omarket, oengine.broker_for("paper", omarket), oclock.now, "2026-09-02",
                                              ostorage.open_trades("paper")[0], {"source": "gemini"}, osettings.all())
    check("일괄 매도 고정이면 AI가 일부 매도를 골라도 실행하지 않음", blocked is False and "일괄 매도" in why and not ostorage.closed_trades("paper"), why[:40])
    set_price(omarket, "BTCUSDT", o["tp_price"] * 1.001)
    tick(oengine, oclock)
    done = ostorage.closed_trades("paper")
    check("일괄 매도: 익절가에서 전량을 한 번에 정리", len(done) == 1 and done[0]["qty"] == o["qty"] and done[0]["exit_reason"] == "익절")

    # --- 수수료를 빼면 남지 않는 자리, 나눌 수 없는 작은 수량
    fsettings, fstorage, fengine, fclock, fmarket = build("exit_fee", {"exit_mode": "split", "exit_split_count": 5})
    fengine.manual_order("BTCUSDT", "long", 200, 5, tp={"unit": "pct", "value": 3}, sl={"unit": "pct", "value": 10}, apply=True)
    fengine.step()
    f = fstorage.open_trades("paper")[0]
    first = f["entry_price"] + (f["tp_price"] - f["entry_price"]) * 0.22
    move = (first / f["entry_price"] - 1) * 100
    set_price(fmarket, "BTCUSDT", first)
    tick(fengine, fclock)
    check("구간에 닿아도 수수료를 빼면 남는 것이 거의 없는 자리에서는 팔지 않음", move < 0.165 and not fstorage.closed_trades("paper"),
          f"가격 변동 {move:.3f}% (왕복 수수료 0.11%의 1.5배 미만)")
    price_now = fmarket.get_price("ETHUSDT")
    tiny = fengine.manual_order("ETHUSDT", "long", price_now * 0.02 / 5 * 1.01, 5, tp={"unit": "pct", "value": 20}, sl={"unit": "pct", "value": 10}, apply=True)
    fengine.step()
    small = next(t for t in fstorage.open_trades("paper") if t["symbol"] == "ETHUSDT")
    set_price(fmarket, "ETHUSDT", small["entry_price"] + (small["tp_price"] - small["entry_price"]) * 0.5)
    tick(fengine, fclock)
    tick(fengine, fclock)
    small_after = next(t for t in fstorage.open_trades("paper") if t["symbol"] == "ETHUSDT")
    notes = [r for r in fstorage.recent_analysis("paper", 50) if r["action"] == "reduce" and r["symbol"] == "ETHUSDT"]
    check("수량이 작아 나눌 수 없으면 한 번에 정리하기로 바꾸고 한 번만 알림", small_after["qty"] == small["qty"] and len(notes) == 1
          and notes[0]["executed"] == 0 and "나눠 팔 수 없" in notes[0]["note"]
          and _json.loads(small_after["context"])["exit"].get("merged") is True, f"수량 {small['qty']:g} (주문 단위 {tiny['qty_step']:g})")

    # --- AI 자동: 진입할 때 AI가 매도 방식을 고르고, 보유 중에는 '일부만 정리'를 고를 수 있다
    def ai_run(name, analyzer, changes=None):
        settings, storage, engine, clock, market = build(name, {"trading_enabled": True, "exit_mode": "auto", "exit_split_count": 3,
                                                                "cross_check": False, "max_open_positions": 1,
                                                                "trade_symbols": ["BTCUSDT"], **(changes or {})}, analyzer=analyzer)
        for _ in range(40):
            if storage.open_trades("paper"):
                break
            clock.now += timedelta(minutes=15)
            engine.step()
        return settings, storage, engine, clock, market

    asettings, astorage, aengine, aclock, amarket = ai_run("exit_ai_split", ExitAI(way="split"))
    opened = astorage.open_trades("paper")
    if not opened:
        check("AI 자동: 점검용 포지션 만들기", False, "포지션이 열리지 않았습니다")
    else:
        a = opened[0]
        entry_note = next((r["note"] for r in astorage.recent_analysis("paper", 20, executed_only=True) if r["action"] in ("long", "short")), "")
        check("AI 자동: AI가 분할 매도를 고르면 그 계획으로 진입하고 이유에 적음", _json.loads(a["context"])["exit"] == {"total": 3, "done": 0}
              and "분할 매도 3회 (AI가 고름)" in entry_note, entry_note[-70:])
        aengine.analyzer.held_action = "reduce"
        aclock.now += timedelta(minutes=15)
        aengine.step()
        sold = astorage.closed_trades("paper")
        check("AI가 '일부만 정리'를 고르면 다음 회차를 바로 팖", len(sold) == 1 and sold[0]["exit_reason"] == "분할 매도 1/3 (AI 판단)"
              and astorage.open_trades("paper") and astorage.open_trades("paper")[0]["qty"] < a["qty"], sold[0]["exit_reason"] if sold else "팔리지 않음")
        aclock.now += timedelta(minutes=15)
        aengine.step()
        aclock.now += timedelta(minutes=15)
        aengine.step()
        sold = [t for t in astorage.closed_trades("paper") if is_part(t)]
        last = next((r for r in astorage.recent_analysis("paper", 5) if r["action"] == "reduce"), None)
        still = astorage.open_trades("paper")
        check("마지막 회차는 일부 매도로 팔지 않음 (익절가, 손절가, 전량 청산으로 정리)", len(sold) == 2 and (not still or (last and last["executed"] == 0
              and "마지막 회차" in last["note"])), last["note"][:40] if last else "")

    bsettings, bstorage, bengine, bclock, bmarket = ai_run("exit_ai_single", ExitAI(way="single"))
    opened = bstorage.open_trades("paper")
    if opened:
        b = opened[0]
        bnote = next((r["note"] for r in bstorage.recent_analysis("paper", 20, executed_only=True) if r["action"] in ("long", "short")), "")
        check("AI 자동: AI가 일괄 매도를 고르면 중간에 팔 계획 없이 진입", _json.loads(b["context"])["exit"] == {"total": 1, "done": 0}
              and "일괄 매도 (AI가 고름)" in bnote and bengine.get_state()["positions"][0]["exit_total"] is None)
        bengine.analyzer.held_action = "reduce"
        bclock.now += timedelta(minutes=15)
        bengine.step()
        sold = bstorage.closed_trades("paper")
        check("AI 자동: 일괄 매도로 들어간 포지션도 AI가 나중에 나눠 팔기로 바꿀 수 있음", len(sold) == 1
              and sold[0]["exit_reason"] == "분할 매도 1/3 (AI 판단)", sold[0]["exit_reason"] if sold else "팔리지 않음")
    else:
        check("AI 자동: 점검용 포지션 만들기 (일괄 매도)", False, "포지션이 열리지 않았습니다")

    csettings, cstorage, cengine, cclock, cmarket = ai_run("exit_ai_add", ExitAI(way="split", entry="split"), {"entry_mode": "auto", "split_count": 3})
    opened = cstorage.open_trades("paper")
    if opened:
        cengine.analyzer.held_action = "reduce"
        cclock.now += timedelta(minutes=15)
        cengine.step()
        cengine.analyzer.held_action = "add"
        cclock.now += timedelta(minutes=15)
        cengine.step()
        last = next((r for r in cstorage.recent_analysis("paper", 5) if r["action"] == "add"), None)
        check("나눠 팔기 시작한 포지션에는 분할 매수를 이어 가지 않음", last is not None and last["executed"] == 0 and "분할 매도를 시작" in last["note"],
              last["note"] if last else "추가 매수 판단이 없음")
    else:
        check("점검용 포지션 만들기 (분할 매수 + 분할 매도)", False, "포지션이 열리지 않았습니다")

    # --- AI가 없을 때는 규칙 전략이 매도 방식을 고른다
    _, rstorage, _, _ = simulate(folder, "exit_rule", 2, {"exit_mode": "auto", "max_hold_minutes": 600})
    plans = [(_json.loads(t["context"]).get("exit") or {}).get("total") for t in rstorage.closed_trades("paper") + rstorage.open_trades("paper")
             if not is_part(t)]
    check("AI가 없으면 규칙 전략이 거래마다 일괄 매도와 분할 매도 중에서 고름", bool(plans) and all(n in (1, 3) for n in plans),
          f"일괄 {plans.count(1)}건, 분할 {plans.count(3)}건")
    piece_rows = [t for t in rstorage.closed_trades("paper") if is_part(t)]
    check("규칙 전략으로 돌려도 분할 매도 기록은 '분할 매도 k/N'으로 남음", all(t["exit_reason"].startswith("분할 매도 ") for t in piece_rows),
          f"분할 매도 기록 {len(piece_rows)}건")
    held_fake = {"symbol": "BTCUSDT", "side": "long", "entry_price": 1.0, "context": _json.dumps({"exit": {"total": 3, "done": 0}})}
    snaps = all_snapshots(SyntheticMarket(seed=3, history_bars=600), ["BTCUSDT"])
    ruled = rule_decide(snaps, {"indices": {}, "available": False}, [dict(held_fake)])[0]
    check("규칙 전략: 보유 포지션에는 유지, 청산, 일부 매도, 추가 매수 중 하나를 돌려줌", ruled["action"] in ("hold", "close", "reduce", "add"), ruled["action"])

    # --- Gemini 프롬프트와 응답 읽기
    gsettings = Settings(folder / "exit_prompt.json")
    analyzer = _gemini.GeminiAnalyzer(gsettings)

    class StubClient:
        model, backup, available = "x", "auto", True

        def __init__(self):
            self.prompt = ""

        def generate_json(self, prompt):
            self.prompt = prompt
            return {"decisions": [
                {"symbol": "BTCUSDT", "action": "reduce", "confidence": 0.6, "reason": "힘이 약해짐"},
                {"symbol": "ETHUSDT", "action": "long", "entry": "single", "exit": "SPLIT", "p_long": 0.8, "p_short": 0.2,
                 "confidence": 0.8, "leverage": 5, "sl_pct": 0.5, "tp_pct": 1.5, "reason": "추세"},
                {"symbol": "XRPUSDT", "action": "short", "exit": "아무거나", "p_long": 0.2, "p_short": 0.8, "confidence": 0.8, "reason": "하락"}]}

    stub = StubClient()
    analyzer.client = stub
    held_row = {"symbol": "BTCUSDT", "side": "long", "entry_price": 100.0, "leverage": 5, "sl_price": 99.0, "tp_price": 103.0,
                "context": _json.dumps({"exit": {"total": 3, "done": 1}})}
    got = {d["symbol"]: d for d in analyzer.decide({"indices": {}}, {}, None, [], [held_row], {"pnl_pct": 0.0, "trades_done": 0, "loss_room": 25.0})}
    check("Gemini 응답에서 매도 방식(exit)과 일부 매도(reduce)를 읽음 (이상한 값은 버림)", got["BTCUSDT"]["action"] == "reduce"
          and got["ETHUSDT"]["exit"] == "split" and got["XRPUSDT"]["exit"] is None)
    check("Gemini에게 매도 방식 규칙과 포지션의 매도 회차를 알려 줌", "- 매도 방식: 네가 선택" in stub.prompt and '"exit_sold": 1' in stub.prompt
          and '"exit_total": 3' in stub.prompt and "reduce" in stub.prompt)
    gsettings.update({"exit_mode": "single"})
    analyzer.decide({"indices": {}}, {}, None, [], [held_row], {"pnl_pct": 0.0, "trades_done": 0, "loss_room": 25.0})
    check("일괄 매도 고정이면 Gemini에게도 고정이라고 알려 줌", "- 매도 방식: 일괄 매도로 고정" in stub.prompt and '"exit_total": 1' in stub.prompt)

    # --- 실제 주문 브로커: 일부만 reduce-only 시장가로 정리하고, 그 주문의 청산 기록만 읽는다
    class PartClient:
        env = "demo"

        def __init__(self):
            self.orders, self.reads = [], 0

        def get_fee_rates(self, symbol):
            return {"taker": 0.00055, "maker": 0.0002}

        def get_last_closed_pnl(self, symbol):
            self.reads += 1
            old = {"closed_pnl": 9.9, "entry_price": 100.0, "exit_price": 90.0, "open_fee": 0.1, "close_fee": 0.1, "qty": 5.0, "updated_ms": 1000}
            new = {"closed_pnl": 1.5, "entry_price": 100.0, "exit_price": 102.0, "open_fee": 0.05, "close_fee": 0.05, "qty": 1.0, "updated_ms": 2000}
            return new if self.orders and self.reads >= 3 else old       # 주문 직후 한 번은 앞 회차의 기록이 그대로 보인다

        def place_market_order(self, symbol, side, qty, stop_loss=None, take_profit=None, reduce_only=False):
            self.orders.append({"symbol": symbol, "side": side, "qty": qty, "reduce_only": reduce_only, "stop_loss": stop_loss})

    real_sleep, _broker.time.sleep = _broker.time.sleep, lambda seconds: None
    try:
        client = PartClient()
        live = LiveBroker(client, None, gsettings)
        closure = live.reduce_position({"symbol": "BTCUSDT", "side": "long", "qty": 1.0, "entry_price": 100.0}, "분할 매도 1/3 (AI 판단)",
                                       {"qty_step": 0.001, "tick_size": 0.1})
    finally:
        _broker.time.sleep = real_sleep
    check("실제 주문: 팔 몫만 reduce-only 시장가로 내고 손절/익절은 건드리지 않음", client.orders == [
        {"symbol": "BTCUSDT", "side": "Sell", "qty": "1.000", "reduce_only": True, "stop_loss": None}], str(client.orders)[:80])
    check("실제 주문: 앞 회차가 아니라 방금 낸 주문의 청산 기록으로 손익을 적음", closure["net_pnl"] == 1.5 and closure["exit_price"] == 102.0
          and closure["reason"] == "분할 매도 1/3 (AI 판단)", f"순손익 {closure['net_pnl']:+.2f}, 청산가 {closure['exit_price']}")

    # --- 화면에서 쓰는 설정 저장
    eapp = create_app(sengine, ssettings, sstorage).test_client()
    saved = eapp.post("/api/settings", json={"exit_mode": "auto", "exit_split_count": 4}, headers=host).get_json()
    state = eapp.get("/api/state", headers=host).get_json()["state"]
    check("매매 설정에서 매도 방식과 분할 횟수를 저장하면 운용 상태에 반영", saved["ok"] and state["exit_mode"] == "auto" and state["exit_split_count"] == 4)


def main() -> int:
    print("Myself Trade 자가 점검을 시작합니다.\n")
    folder = Path(tempfile.mkdtemp(prefix="myself_trade_test_"))

    print("1. 수수료와 손익 계산")
    pnl = calc_trade_pnl("long", 100.0, 101.0, 10.0, 0.00055, 0.00055)
    check("롱 손익 = (101-100) x 10 - 수수료", abs(pnl["net_pnl"] - (10 - 0.55 - 0.5555)) < 1e-9, f"순손익 {pnl['net_pnl']:.4f}")
    pnl = calc_trade_pnl("short", 100.0, 101.0, 10.0, 0.00055, 0.00055)
    check("숏은 가격이 오르면 손실", pnl["net_pnl"] < -10)
    check("수량을 거래소 단위에 맞춰 내림", round_step(0.123456, 0.001) == 0.123 and round_step(157.9, 1) == 157.0)
    sample = {"chart": {"result": [{"meta": {"regularMarketPrice": 110.0},
                                    "indicators": {"quote": [{"close": [100.0, None, 105.0, 110.0]}]}}]}}
    parsed = parse_yahoo_chart(sample)
    check("미국 증시 데이터 해석", parsed is not None and parsed["chg_1d"] == round((110 / 105 - 1) * 100, 2))

    print("\n2. 가상 시세로 3일 매매")
    settings, storage, engine, clock = simulate(folder, "normal", 3, {"invest_pct": 20, "leverage_mode": "manual",
                                                                    "manual_leverage": 5})
    trades = storage.closed_trades("paper")
    check("거래가 실제로 체결됨", len(trades) > 0, f"완료 {len(trades)}회")
    consistent = all(abs((t["gross_pnl"] - t["fee"]) - t["net_pnl"]) < 1e-6 for t in trades)
    check("모든 거래에서 순손익 = 손익 - 수수료", consistent)
    check("수수료가 모든 거래에 반영됨", all(t["fee"] > 0 for t in trades))
    check("레버리지 수동 고정값(5배) 적용", all(t["leverage"] == 5 for t in trades))
    open_trades = storage.open_trades("paper")
    open_fees = sum(t["fee"] for t in open_trades)
    balance = storage.kv_get("paper_balance")
    expected = 1000.0 + sum(t["net_pnl"] for t in trades) - open_fees
    check("잔고 = 시작 잔고 + 순손익 합계", abs(balance - expected) < 1e-6, f"잔고 {balance:.4f} USDT")
    per_day: dict[str, int] = {}
    for t in trades:
        per_day[t["day"]] = per_day.get(t["day"], 0) + 1
    check("하루 최대 거래 횟수 이내", all(n <= settings.get("max_trades_per_day") for n in per_day.values()), str(per_day))
    check("포지션 하나의 투자 금액이 오늘 원금의 10% (20% / 2포지션) 이내", all(t["invest_pct"] <= 10.05 for t in trades),
          f"가장 큰 값 {max(t['invest_pct'] for t in trades):.2f}%")
    from risk_manager import RiskManager
    rm = RiskManager(settings)          # 증거금 투자율 20%, 동시 보유 2개 -> 포지션 하나에 오늘 원금의 10%
    unit = {"qty_step": 0.001, "min_qty": 0.001, "max_qty": 0, "min_notional": 5.0, "tick_size": 0.01}
    common = dict(start_equity=1000.0, price=100.0, leverage=5, sl_pct=1.0, fee_rate=0.00055, instrument=unit)
    after_gain = rm.position_size(equity=1300.0, pnl_pct=30.0, **common)
    after_loss = rm.position_size(equity=900.0, pnl_pct=-10.0, **common)
    check("투자 금액은 오늘 원금 기준 (하루 중 수익이나 손실이 나도 같음)",
          abs(after_gain["margin"] - 100.0) < 0.05 and abs(after_loss["margin"] - 100.0) < 0.05,
          f"원금 1000, 수익 후 {after_gain['margin']:.2f} / 손실 후 {after_loss['margin']:.2f} USDT")
    short = rm.position_size(equity=150.0, pnl_pct=-5.0, used_margin=100.0, **common)
    check("남은 자산보다 많이 투자하지 않음", 0 < short["margin"] <= 50.0, f"남은 자산 50 -> {short['margin']:.2f} USDT")
    rm_full = Settings(folder / "limit.json")
    rm_full.update({"invest_pct": 100, "max_open_positions": 4})
    full = RiskManager(rm_full)          # 투자율 100%, 포지션 4개 -> 하나에 250, 합계 한도 1000
    fourth = full.position_size(equity=1003.0, pnl_pct=0.3, used_margin=496.90 + 248.10 + 257.77, **common)
    check("이미 투자한 금액과 합쳐 투자 한도를 넘지 않음", fourth["qty"] == 0 and "투자 한도" in fourth["note"],
          "497 + 248 + 258 = 1003을 쓰는 중이면 네 번째는 사지 않음")
    third = full.position_size(equity=1003.0, pnl_pct=0.3, used_margin=496.90 + 248.10, **common)
    check("남은 한도만큼만 투자", 0 < third["margin"] <= 1000.0 - 745.0 + 0.01, f"남은 한도 255 -> {third['margin']:.2f} USDT")
    msettings2 = Settings(folder / "limit_run.json")
    msettings2.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "invest_pct": 100, "max_open_positions": 2,
                       "leverage_mode": "manual", "manual_leverage": 5, "max_hold_minutes": 1440, "min_confidence": 0.3,
                       "entry_mode": "single"})
    mstorage2 = Storage(folder / "limit_run.db")
    mstart = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
    mclock2 = Clock(mstart)
    mmarket2 = SyntheticMarket(seed=9, history_bars=3000, start_ms=int(mstart.timestamp() * 1000))
    mengine2 = TradeEngine(msettings2, mstorage2, market=mmarket2, clock=mclock2, analyzer=NoAI(), macro=NoMacro())
    worst, most_open = 0.0, 0
    for index in range(2 * 288):
        if index == 40:
            msettings2.update({"max_open_positions": 4})      # 포지션을 들고 있는 중간에 설정을 바꾼다
        mengine2.step(); mmarket2.advance(1); mclock2.now += timedelta(minutes=5)
        opened = mstorage2.open_trades("paper")
        day_row = mstorage2.get_day("paper", mengine2.get_state()["day"])
        if opened and day_row:
            worst = max(worst, sum(t["margin"] for t in opened) / day_row["start_equity"])
            most_open = max(most_open, len(opened))
    check("포지션 수 설정을 중간에 바꿔도 투자금 합계가 오늘 원금을 넘지 않음", 0 < worst <= 1.002 and most_open >= 3,
          f"가장 많이 투자했을 때 원금의 {worst * 100:.1f}%, 동시에 최대 {most_open}개 보유")
    day_rows = {d["day"]: d for d in storage.days_between("paper", "2026-01-01", "2027-01-01")}
    check("다음 날 원금 = 전날 결과 금액", len(day_rows) >= 2 and all(
        abs(day_rows[a]["end_equity"] - day_rows[b]["start_equity"]) < 1e-6
        for a, b in zip(sorted(day_rows), sorted(day_rows)[1:])),
          ", ".join(f"{d[5:]}: {day_rows[d]['start_equity']:.2f}" for d in sorted(day_rows)))
    days = storage.days_between("paper", "2026-01-01", "2027-01-01")
    check("날짜가 바뀌면 전날 회고가 남음", sum(1 for d in days if d["review"]) >= 2, f"기록된 날 {len(days)}일")
    check("분석(투자 이유) 기록이 남음", len(storage.recent_analysis("paper", 10, executed_only=True)) > 0)

    print("\n3. 하루 거래 횟수 제한")
    limited, lstorage, _, _ = simulate(folder, "limit", 2, {"max_trades_per_day": 3, "min_trades_per_day": 1,
                                                          "invest_pct": 10, "max_hold_minutes": 30})
    counts: dict[str, int] = {}
    for t in lstorage.closed_trades("paper"):
        if not is_part(t):                 # 분할 매도로 나눠 판 포지션은 전부 정리된 때 한 번으로 센다
            counts[t["day"]] = counts.get(t["day"], 0) + 1
    check("하루 3회로 제한하면 3회를 넘지 않음 (분할 매도로 나눠 판 포지션은 한 번으로 셈)",
          bool(counts) and all(n <= 3 for n in counts.values()), str(counts))

    print("\n4. 일일 손실 제한 (손실 록)")
    locked_days = []
    for lock_seed in (17, 19, 20, 14, 21, 11):     # 손실이 한도에 닿는 날이 나오는 가상 시세를 찾는다
        risky, rstorage, rengine, rclock = simulate(folder, f"lock{lock_seed}", 2, {
            "invest_pct": 100, "leverage_mode": "manual", "manual_leverage": 50, "daily_loss_limit_pct": 3,
            "max_open_positions": 4, "min_confidence": 0.3, "max_hold_minutes": 600}, seed=lock_seed)
        locked_days = [d for d in rstorage.days_between("paper", "2026-01-01", "2027-01-01") if d["locked"]]
        if locked_days:
            break
    check("손실 한도에 닿으면 그날 잠김", len(locked_days) > 0, f"잠긴 날 {len(locked_days)}일")
    if locked_days:
        day = locked_days[0]
        day_trades = rstorage.trades_on_day("paper", day["day"])
        forced = [t["exit_time"] for t in day_trades if t["exit_reason"] == "일일 손실 제한"]
        if forced:
            lock_time = max(forced)
        else:      # 손절로 이미 다 정리된 뒤에 잠겼다면, 누적 손실이 한도를 넘긴 거래의 청산 시각이 잠긴 시각이다
            running, lock_time = 0.0, day_trades[-1]["exit_time"]
            for t in day_trades:
                running += t["net_pnl"]
                if running <= -0.03 * day["start_equity"]:
                    lock_time = t["exit_time"]
                    break
        late = [t for t in rstorage.closed_trades("paper") if t["entry_time"] > lock_time and t["day"] == day["day"]]
        check("잠긴 뒤에는 그날 새 진입이 없음", not late)
        worst = [min(p["pnl_pct"] for p in rstorage.equity_series("paper", day=d["day"])) for d in locked_days]
        check("잠긴 날의 손실이 한도 근처에서 멈춤", min(worst) > -3 * 2,
              "잠긴 날 최저 수익률 " + ", ".join(f"{v:.2f}%" for v in worst) + " (한도 -3%)")
    else:
        check("잠긴 뒤에는 그날 새 진입이 없음", False, "잠긴 날이 없어 확인 불가")

    print("\n5. 리포트와 엑셀")
    for period in ("1d", "1w", "1m", "1y"):
        report = build_report(storage, "paper", period, "Asia/Seoul", now=clock.now)
        check(f"{period} 리포트", report["summary"]["trades"] >= 0 and isinstance(report["series"], list),
              f"거래 {report['summary']['trades']}회, 수익률 {report['summary']['return_pct']:+.2f}%")
    book = load_workbook(export_xlsx(storage, "paper", "1m", "Asia/Seoul", now=clock.now))
    check("엑셀 파일 생성", {"요약", "거래내역", "일별수익", "코인별"} <= set(book.sheetnames), ", ".join(book.sheetnames))
    check("엑셀 거래내역 행 수 = 거래 수", book["거래내역"].max_row - 1 == len(trades))

    print("\n6. 대시보드")
    app = create_app(engine, settings, storage)
    client = app.test_client()
    for path in ("/", "/api/state", "/api/report?period=1w", "/api/returns", "/api/analysis", "/api/settings",
                 "/api/price?symbol=BTCUSDT", "/api/export?period=1w"):
        response = client.get(path, headers={"Host": "127.0.0.1:8000"})
        check(f"GET {path}", response.status_code == 200)
    engine.step()
    repeated = [client.get("/api/state", headers={"Host": "127.0.0.1:8000"}).status_code for _ in range(3)]
    check("상태 조회를 반복해도 정상", repeated == [200, 200, 200], f"보유 포지션 {len(storage.open_trades('paper'))}개")
    response = client.post("/api/toggle", json={"enabled": False}, headers={"Host": "127.0.0.1:8000"})
    check("자동 매매 Off 전환", response.status_code == 200 and settings.get("trading_enabled") is False)
    response = client.post("/api/settings", json={"mode": "live", "confirm_live": True}, headers={"Host": "127.0.0.1:8000"})
    import config as _config
    if not _config.has_bybit_keys():
        check("API 키 없이는 실전 모드로 바꿀 수 없음", response.status_code == 400 and settings.get("mode") == "paper")
    response = client.post("/api/keys", json={"bybit_key": "한쪽만"}, headers={"Host": "127.0.0.1:8000"})
    if not _config.has_bybit_keys():
        check("Bybit 키는 Key와 Secret을 함께 넣어야 저장됨", response.status_code == 400)
    host = {"Host": "127.0.0.1:8000"}
    info = client.get("/api/paper", headers=host).get_json()
    check("모의투자 요약 조회", info["ok"] and info["trades"] == len(storage.closed_trades("paper")),
          f"가상 잔고 {info['balance']:.2f} USDT, 누적 {info['return_pct']:+.2f}%")
    check("가상 잔고 요약의 시작 잔고가 1000 USDT", abs(info["base"] - 1000.0) < 1e-6)
    if storage.open_trades("paper"):
        blocked = client.post("/api/settings", json={"mode": "demo"}, headers=host)
        check("보유 포지션이 있으면 모드를 바꿀 수 없음", blocked.status_code == 400)
    client.post("/api/paper_reset", json={"start_balance": 500}, headers=host)
    info = client.get("/api/paper", headers=host).get_json()
    check("모의투자 다시 시작", info["trades"] == 0 and info["balance"] == 500 and settings.get("paper_start_balance") == 500)
    if not _config.has_bybit_keys():
        off = client.post("/api/settings", json={"mode": "demo"}, headers=host)
        check("Bybit 키가 없으면 모의투자 모드를 끌 수 없음", off.status_code == 400 and settings.get("mode") == "paper")
    response = client.get("/api/state", headers={"Host": "evil.example.com"})
    check("다른 주소에서 온 요청은 거부", response.status_code == 403)

    print("\n참고: 위 수익률은 무작위로 만든 가상 시세의 결과입니다. 실제 시장 성과를 뜻하지 않습니다.")
    print("\n7. 새 기능 (차트, 매매 방향, 개별 청산, 학습, 메일)")
    host = {"Host": "127.0.0.1:8000"}
    from config import TIMEFRAMES
    sizes = {}
    for value, name, _ in TIMEFRAMES:
        body = client.get(f"/api/price?symbol=BTCUSDT&tf={value}", headers=host).get_json()
        sizes[name] = len(body.get("candles", [])) if body.get("ok") else -1
    check("타임프레임 13종 모두 캔들 조회", len(sizes) == 13 and all(n > 0 for n in sizes.values()), str(sizes))
    body = client.get("/api/price?symbol=BTCUSDT&tf=180", headers=host).get_json()
    c = body["candles"][-2]
    check("캔들에 시가/고가/저가/종가가 있고 고가 >= 저가", c["h"] >= max(c["o"], c["c"]) and c["l"] <= min(c["o"], c["c"]))
    latest = client.get("/api/price?symbol=BTCUSDT&tf=15&count=300", headers=host).get_json()
    oldest_t = latest["candles"][0]["t"]
    older = client.get(f"/api/price?symbol=BTCUSDT&tf=15&count=300&end={oldest_t - 1}", headers=host).get_json()
    check("과거 구간 이어서 불러오기 (앞 구간과 겹치지 않고 시간순)", older["ok"] and len(older["candles"]) == 300
          and older["candles"][-1]["t"] < oldest_t and older["has_more"] is True
          and all(a["t"] < b["t"] for a, b in zip(older["candles"], older["candles"][1:])))
    pages, cursor, total = 0, oldest_t, len(latest["candles"])
    while pages < 30:
        page = client.get(f"/api/price?symbol=BTCUSDT&tf=15&count=1000&end={cursor - 1}", headers=host).get_json()
        pages += 1
        if not page["candles"]:
            break
        total += len(page["candles"]); cursor = page["candles"][0]["t"]
        if not page["has_more"]:
            break
    check("가장 오래된 시세(상장 시점)까지 닿으면 더 없음으로 표시", page["has_more"] is False and pages < 30, f"15분봉 {total}개, {pages}번에 나눠 조회")
    ticker = client.get("/api/ticker?symbol=ETHUSDT", headers=host).get_json()
    check("현재가 조회", ticker["ok"] and ticker["price"] > 0, ticker["feed"]["label"])

    for direction in ("long", "short"):
        _, dstorage, _, _ = simulate(folder, "dir_" + direction, 1.5, {"direction_mode": direction})
        sides = {t["side"] for t in dstorage.closed_trades("paper")} | {t["side"] for t in dstorage.open_trades("paper")}
        check(f"매매 방향 '{'롱만' if direction == 'long' else '숏만'}' 설정 시 그 방향만 진입", sides <= {direction} and len(sides) == 1, str(sides))

    msettings, mstorage, mengine, mclock = simulate(folder, "manual", 0.6, {"max_hold_minutes": 1440})
    opened = mstorage.open_trades("paper")
    if opened:
        mapp = create_app(mengine, msettings, mstorage).test_client()
        target = opened[0]
        answer = mapp.post("/api/close_position", json={"id": target["id"]}, headers=host)
        mengine.step()
        closed = [t for t in mstorage.closed_trades("paper") if t["id"] == target["id"]]
        check("개별 청산 버튼으로 한 포지션만 청산", answer.status_code == 200 and closed and closed[0]["exit_reason"] == "수동 청산"
              and len(mstorage.open_trades("paper")) == len(opened) - 1)
        state = mapp.get("/api/state", headers=host).get_json()["state"]
        check("보유 포지션에 투자 금액(증거금) 표시", all(p["margin"] > 0 and p["notional"] > p["margin"] for p in state["positions"]) or not state["positions"])
    else:
        check("개별 청산 버튼으로 한 포지션만 청산", False, "보유 포지션이 없어 확인 불가")

    nsettings = Settings(folder / "net.json")
    nsettings.update({"mode": "paper", "trading_enabled": True, "timezone": "Asia/Seoul", "invest_pct": 100,
                      "leverage_mode": "manual", "manual_leverage": 26, "max_hold_minutes": 1440, "min_confidence": 0.3})
    nstorage = Storage(folder / "net.db")
    nstart = datetime(2026, 9, 1, 15, 0, tzinfo=timezone.utc)
    nclock = Clock(nstart)
    nmarket = SyntheticMarket(seed=3, history_bars=3000, start_ms=int(nstart.timestamp() * 1000))
    nengine = TradeEngine(nsettings, nstorage, market=nmarket, clock=nclock, analyzer=NoAI(), macro=NoMacro())
    state = {}
    for _ in range(200):
        nengine.step(); nmarket.advance(1); nclock.now += timedelta(minutes=5)
        state = nengine.get_state()
        if state["positions"] and nclock.now.hour != 15:
            nengine.step(); state = nengine.get_state()
            break
    if state.get("positions"):
        ps = state["positions"]
        realized = sum(t["net_pnl"] for t in nstorage.trades_on_day("paper", state["day"]))
        open_net = sum(p["unrealized"] for p in ps)
        check("포지션 순손익 = 가격 손익 - 진입 수수료 - 청산 수수료(예상)",
              all(abs(p["unrealized"] - (p["gross_pnl"] - p["fee_paid"] - p["fee_exit"])) < 1e-9 and p["fee_paid"] > 0 and p["fee_exit"] > 0 for p in ps))
        check("오늘 수익 = 청산 완료 거래의 순손익 + 보유 포지션 순손익 합계", abs(state["pnl_amount"] - (realized + open_net)) < 1e-6,
              f"{state['pnl_amount']:+.4f} = {realized:+.4f} + {open_net:+.4f}")
        check("화면의 '청산 완료 + 보유 중' 값이 오늘 수익과 일치", abs(state["closed_net"] + state["open_net"] - state["pnl_amount"]) < 1e-9
              and abs(state["closed_net"] - realized) < 1e-6)
        check("순자산 = 거래소 기준 자산 - 청산 수수료 예상", abs(state["equity"] - (state["gross_equity"] - sum(p["fee_exit"] for p in ps))) < 1e-9)
        for trade in nstorage.open_trades("paper"):
            nengine.request_close(trade["id"])
        before_close = nengine.get_state()["equity"]
        nengine.step()
        after_close = nengine.get_state()
        check("전부 청산해도 순자산이 그대로 (청산 수수료를 미리 반영)", not after_close["positions"]
              and abs(after_close["equity"] - before_close) < before_close * 0.0005,
              f"청산 전 {before_close:.4f} -> 청산 후 {after_close['equity']:.4f}")
    else:
        check("포지션 순손익 = 가격 손익 - 진입 수수료 - 청산 수수료(예상)", False, "보유 포지션이 없어 확인 불가")

    import json as _json
    _, bstorage, bengine, _ = simulate(folder, "btc_only", 2, {"trade_symbols": ["BTCUSDT"], "entry_mode": "split",
                                                                 "split_count": 3, "max_hold_minutes": 600})
    every = [t for t in bstorage.closed_trades("paper") + bstorage.open_trades("paper") if not is_part(t)]
    check("거래할 코인으로 비트코인만 고르면 비트코인만 거래", bool(every) and {t["symbol"] for t in every} == {"BTCUSDT"}, f"{len(every)}건")
    splits = [_json.loads(t["context"]).get("split") for t in every]
    check("분할 매수 고정: 모든 진입이 3회로 나뉨", all(sp and sp["total"] == 3 for sp in splits))
    grown = [t for t, sp in zip(every, splits) if sp and sp["done"] > 1]
    check("분할 매수: 방향이 유지되면 다음 회차를 이어서 삼", len(grown) > 0, f"{len(every)}건 중 {len(grown)}건이 2회차 이상")
    check("추가 매수한 이유가 분석 기록에 남음", any(r["action"] == "add" for r in bstorage.recent_analysis("paper", 500, executed_only=True)))
    if grown:
        t = grown[0]; ctx = _json.loads(t["context"])
        sl_gap = abs(t["entry_price"] - t["sl_price"]) / t["entry_price"] * 100
        check("추가 매수 뒤 손절가를 새 평균 단가 기준으로 다시 잡음", abs(sl_gap - ctx["sl_pct"]) < 0.05,
              f"평균 단가 {t['entry_price']:.2f}, 손절 폭 {sl_gap:.3f}% (설정 {ctx['sl_pct']}%)")
        check("추가 매수한 수량이 한 회차 수량 x 산 횟수를 넘지 않음", t["qty"] <= ctx["split"]["unit_qty"] * ctx["split"]["done"] + 1e-9)
        check("분할 매수를 해도 투자 금액이 오늘 원금 x 투자율을 넘지 않음", all(x["invest_pct"] <= 10.05 for x in every),
              f"가장 큰 값 {max(x['invest_pct'] for x in every):.2f}% (설정 10%)")
    closed_b = bstorage.closed_trades("paper")
    expected_b = 1000.0 + sum(t["net_pnl"] for t in closed_b) - sum(t["fee"] for t in bstorage.open_trades("paper"))
    check("분할 매수에서도 잔고 = 시작 잔고 + 순손익 합계", abs(bstorage.kv_get("paper_balance") - expected_b) < 1e-6)
    check("코인을 한 개만 고르면 투자 금액을 그 코인에 모두 씀", bengine.risk.slots() == 1)
    _, sstorage, _, _ = simulate(folder, "single", 1, {"entry_mode": "single"})
    singles = sstorage.closed_trades("paper") + sstorage.open_trades("paper")
    check("일괄 매수 고정: 나눠 사지 않음", bool(singles) and all(not _json.loads(t["context"]).get("split") for t in singles))
    _, astorage, _, _ = simulate(folder, "auto_entry", 2, {"entry_mode": "auto"})
    autos = [bool(_json.loads(t["context"]).get("split")) for t in astorage.closed_trades("paper") + astorage.open_trades("paper")]
    check("AI 자동: 상황에 따라 일괄과 분할을 골라 씀", bool(autos), f"분할 {sum(autos)}건, 일괄 {len(autos) - sum(autos)}건")
    bad_coin = client.post("/api/settings", json={"trade_symbols": []}, headers=host)
    check("거래할 코인을 하나도 고르지 않으면 저장되지 않음", bad_coin.status_code == 400)

    with_context = [t for t in storage.closed_trades("paper")]
    hsettings, hstorage, hengine, hclock = simulate(folder, "history", 2, {})
    history = hengine._history("paper")
    sample = hstorage.closed_trades("paper")[0]
    check("진입 당시 시장 상황 저장", "h1_trend" in _json.loads(sample["context"]))
    check("과거 거래 성과 집계 (코인, 방향별)", bool(history) and all("win_rate" in st for sides in history.values() for st in sides.values()),
          f"{sum(len(v) for v in history.values())}개 조합")

    import mailer
    posted = []          # 간편 발송으로 나간 요청
    confirmed = {"me@naver.com"}   # 주소 확인(Activate)을 마친 주소라고 가정

    class FakeResponse:
        def __init__(self, data, code=200): self._data, self.status_code = data, code
        def json(self): return self._data

    def fake_post(url, json=None, headers=None, timeout=None):
        to = url.rsplit("/", 1)[-1]
        posted.append((to, json))
        if to in confirmed:
            return FakeResponse({"success": "true", "message": "The form was submitted successfully."})
        return FakeResponse({"success": "false", "message": "This form needs Activation. We've sent you an email containing an 'Activate Form' link."})

    real_post, mailer.requests.post = mailer.requests.post, fake_post
    try:
        gsettings, gstorage, gengine, gclock = simulate(folder, "mailtest", 0.5, {})
        gengine.mail_async = False
        gapp = create_app(gengine, gsettings, gstorage).test_client()
        early = gapp.post("/api/mail", json={"mail_enabled": True}, headers=host).get_json()
        check("받는 주소가 없으면 자동 발송이 켜지지 않음", early["warning"] and gsettings.get("mail_enabled") is False)
        for address in ("me@naver.com", "friend@daum.net"):
            gapp.post("/api/mail/recipients", json={"add": address}, headers=host)
        duplicate = gapp.post("/api/mail/recipients", json={"add": "ME@naver.com"}, headers=host)
        bad = gapp.post("/api/mail/recipients", json={"add": "주소아님"}, headers=host)
        check("받는 주소 추가 (중복과 잘못된 형식은 거부)", gsettings.get("mail_recipients") == ["me@naver.com", "friend@daum.net"]
              and duplicate.status_code == 400 and bad.status_code == 400)
        check("비밀번호 없이 주소만으로 발송 준비 완료", mailer.mail_ready(gsettings) == (True, ""))
        tested = gapp.post("/api/mail/test", json={}, headers=host).get_json()
        states = {r["email"]: r["state"] for r in tested["recipients"]}
        check("확인 메일: 확인된 주소는 발송, 새 주소는 확인 대기로 표시",
              states == {"me@naver.com": "ok", "friend@daum.net": "pending"}, str(states))
        gapp.post("/api/mail", json={"mail_enabled": True, "mail_hour": 21}, headers=host)
        before = len(posted)
        for _ in range(int(1.6 * 288)):          # 이틀째 21시를 지나도록 돌린다
            gengine.step(); gengine._fixed_market.advance(1); gclock.now += timedelta(minutes=5)
        reports = [x for x in posted[before:] if "일일 매매 보고서" in x[1]["_subject"]]
        per_address = {to: sum(1 for x in reports if x[0] == to) for to in ("me@naver.com", "friend@daum.net")}
        check("21시가 되면 주소마다 하루에 한 번 보고서 발송", per_address == {"me@naver.com": 2, "friend@daum.net": 2}, str(per_address))
        if reports:
            body = reports[0][1]
            check("보고서에 수익률, 거래, 잘한 점, 실수와 개선점 포함",
                  all(key in body for key in ("하루 수익률", "거래", "주요 거래", "잘한 점", "실수와 개선점", "다음 방침")), body["_subject"])
        removed = gapp.post("/api/mail/recipients", json={"remove": "friend@daum.net"}, headers=host).get_json()
        check("받는 주소 삭제", gsettings.get("mail_recipients") == ["me@naver.com"] and len(removed["recipients"]) == 1)
        preview = gapp.post("/api/mail/preview", json={}, headers=host).get_json()
        check("보고서 미리보기", preview["ok"] and "<table" in preview["html"])

        # 내 메일 계정(SMTP)으로 보내는 방식
        sent = []

        class FakeSMTP:
            def __init__(self, host, port, **kwargs): self.host, self.port = host, port
            def __enter__(self): return self
            def __exit__(self, *args): return False
            def login(self, user, password): sent.append(("login", user, password))
            def send_message(self, message): sent.append(("mail", message))

        mailer.smtplib.SMTP_SSL = FakeSMTP
        old_password = _config.os.environ.get("MAIL_PASSWORD")
        need = gapp.post("/api/mail", json={"mail_method": "smtp"}, headers=host).get_json()
        check("내 메일 계정 방식은 계정과 앱 비밀번호가 있어야 함", need["ready"] is False)
        _config.os.environ["MAIL_PASSWORD"] = "abcdabcdabcdabcd"
        try:
            gsettings.update({"mail_from": "bot@example.com"})
            now_sent = gapp.post("/api/mail/send_now", json={}, headers=host).get_json()
            mails = [x[1] for x in sent if x[0] == "mail"]
            check("내 메일 계정 방식으로 보고서 발송 (HTML 본문 포함)", now_sent["ok"] and len(mails) == 1
                  and mails[0]["To"] == "me@naver.com" and mails[0].get_body(preferencelist=("html",)) is not None)
        finally:
            if old_password is None:
                _config.os.environ.pop("MAIL_PASSWORD", None)
            else:
                _config.os.environ["MAIL_PASSWORD"] = old_password

        old = Settings(folder / "old_mail.json")
        (folder / "old_mail.json").write_text('{"mail_to": "old@gmail.com", "mail_from": "old@gmail.com", "settings_version": 2}', encoding="utf-8")
        old.load()
        check("예전 설정의 받는 주소를 새 목록으로 옮김", old.get("mail_recipients") == ["old@gmail.com"] and old.get("mail_method") == "smtp")
    finally:
        mailer.requests.post = real_post

    section_v9(folder, host, storage, engine, clock, settings)
    section_v10(folder, host)
    section_v12(folder, host)

    failed = [r for r in RESULTS if not r[1]]
    print(f"\n점검 {len(RESULTS)}개 중 {len(RESULTS) - len(failed)}개 통과" + ("" if failed else ". 모두 정상입니다."))
    for name, _, detail in failed:
        print(f"  실패: {name} {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
