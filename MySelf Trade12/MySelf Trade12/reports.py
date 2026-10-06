"""Myself Trade - 성과 리포트 (기간별 통계, 차트용 데이터, 엑셀 내보내기)."""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import APP_NAME, MODES, SYMBOLS, coin_meta
from utils import day_start_utc, fmt_local, iso, local_day, minutes_between, parse_iso, utcnow

PERIODS = {"1d": ("하루", 0), "1w": ("한 주", 6), "1m": ("한 달", 29), "1y": ("한 해", 364)}


def period_start(period: str, tz: str, now=None):
    """기간의 시작 시각(UTC). 1d는 오늘 0시(현지), 1w는 오늘 포함 최근 7일."""
    now = now or utcnow()
    days_back = PERIODS.get(period, PERIODS["1d"])[1]
    return day_start_utc(now, tz) - timedelta(days=days_back)


def trade_view(t: dict, tz: str) -> dict:
    """청산이 끝난 거래 한 건을 화면/엑셀용으로 바꾼다. 진입 시각과 청산 시각을 따로 담는다."""
    return {
        "id": t["id"], "symbol": t["symbol"], "ticker": coin_meta(t["symbol"])["ticker"],
        "name": coin_meta(t["symbol"])["name"],
        "side": t["side"], "leverage": t["leverage"], "qty": t["qty"],
        "entry_price": t["entry_price"], "exit_price": t["exit_price"],
        "entry_time": fmt_local(t["entry_time"], tz, "%m-%d %H:%M"),
        "exit_time": fmt_local(t["exit_time"], tz, "%m-%d %H:%M"),
        "entry_at": fmt_local(t["entry_time"], tz, "%Y-%m-%d %H:%M:%S"),      # 포지션이 체결된 정확한 일시
        "exit_at": fmt_local(t["exit_time"], tz, "%Y-%m-%d %H:%M:%S"),        # 포지션이 끝난 정확한 일시
        "held_minutes": max(0, int(minutes_between(parse_iso(t["entry_time"]), parse_iso(t["exit_time"])))) if t["exit_time"] else None,
        "day": t["day"],
        "margin": t["margin"], "notional": t["entry_price"] * t["qty"], "invest_pct": t["invest_pct"],
        "gross_pnl": t["gross_pnl"], "fee": t["fee"],
        "net_pnl": t["net_pnl"], "return_pct": t["return_pct"], "equity_return_pct": t["equity_return_pct"],
        "exit_reason": t["exit_reason"], "entry_reason": t["entry_reason"], "source": t["source"],
    }


def build_pnl_board(storage, mode: str, tz: str, positions: list[dict], start_equity: float, now=None) -> dict:
    """오늘의 코인별 손익 (수익 코인 / 손실 코인).

    오늘 청산한 거래의 순손익 + 보유 중인 포지션의 순손익(지금 청산 기준)을 코인별로 더한다.
    어제부터 들고 온 포지션은 오늘 0시 이후에 생긴 손익만 센다. 그래서 합계가 대시보드의 '오늘 수익'과 같다.
    """
    now = now or utcnow()
    day = local_day(now, tz)
    carry = storage.kv_get(f"carry_{mode}_{day}") or {}       # 오늘 0시에 들고 있던 포지션의 그때 순손익
    def blank(symbol: str) -> dict:
        meta = coin_meta(symbol)
        return {"symbol": symbol, "name": meta["name"], "ticker": meta["ticker"], "net": 0.0, "margin": 0.0,
                "notional": 0.0, "trades": 0, "open": 0}

    coins = {symbol: blank(symbol) for symbol in list(SYMBOLS)}
    for t in storage.trades_on_day(mode, day):
        coin = coins.setdefault(t["symbol"], blank(t["symbol"]))      # 목록에서 지운 코인의 거래도 빠뜨리지 않는다
        coin["net"] += (t["net_pnl"] or 0.0) - float(carry.get(str(t["id"]), 0.0))
        coin["margin"] += t["margin"] or 0.0
        coin["notional"] += t["entry_price"] * t["qty"]
        coin["trades"] += 1
    for p in positions or []:
        coin = coins.setdefault(p["symbol"], blank(p["symbol"]))
        coin["net"] += (p.get("unrealized") or 0.0) - float(carry.get(str(p["id"]), 0.0))
        coin["margin"] += p.get("margin") or 0.0
        coin["notional"] += p.get("notional") or 0.0
        coin["open"] += 1
    rows = []
    for coin in coins.values():
        coin["return_pct"] = coin["net"] / coin["margin"] * 100 if coin["margin"] else 0.0      # 투자금(증거금) 대비
        coin["equity_pct"] = coin["net"] / start_equity * 100 if start_equity else 0.0          # 오늘 원금 대비
        rows.append(coin)
    profits = sorted((c for c in rows if c["net"] > 1e-9), key=lambda c: c["net"], reverse=True)
    losses = sorted((c for c in rows if c["net"] < -1e-9), key=lambda c: c["net"])
    profit_total, loss_total = sum(c["net"] for c in profits), sum(c["net"] for c in losses)

    def share(value: float) -> float:
        return value / start_equity * 100 if start_equity else 0.0

    return {"day": day, "coins": rows, "profits": profits, "losses": losses, "start_equity": start_equity,
            "profit_total": profit_total, "loss_total": loss_total, "net_total": profit_total + loss_total,
            "profit_pct": share(profit_total), "loss_pct": share(loss_total), "net_pct": share(profit_total + loss_total),
            "active": any(c["trades"] or c["open"] for c in rows)}


def build_report(storage, mode: str, period: str, tz: str, now=None) -> dict:
    now = now or utcnow()
    if period not in PERIODS:
        period = "1d"
    start = period_start(period, tz, now)
    start_day, end_day = local_day(start + timedelta(hours=12), tz), local_day(now, tz)
    trades = storage.closed_trades(mode, iso(start))
    days = storage.days_between(mode, start_day, end_day)

    start_equity = days[0]["start_equity"] if days else 0.0
    net = sum(t["net_pnl"] or 0 for t in trades)
    gross = sum(t["gross_pnl"] or 0 for t in trades)
    fees = sum(t["fee"] or 0 for t in trades)
    wins = [t for t in trades if (t["net_pnl"] or 0) > 0]
    losses = [t for t in trades if (t["net_pnl"] or 0) <= 0]

    # 일별 집계 (거래가 없던 날도 0으로 표시)
    by_day: dict[str, dict] = {d["day"]: {"day": d["day"], "net_pnl": 0.0, "fee": 0.0, "trades": 0,
                                          "start_equity": d["start_equity"]} for d in days}
    for t in trades:
        row = by_day.setdefault(t["day"], {"day": t["day"], "net_pnl": 0.0, "fee": 0.0, "trades": 0,
                                           "start_equity": start_equity})
        row["net_pnl"] += t["net_pnl"] or 0
        row["fee"] += t["fee"] or 0
        row["trades"] += 1
    daily, cumulative = [], 0.0
    for key in sorted(by_day):
        row = by_day[key]
        cumulative += row["net_pnl"]
        row["return_pct"] = row["net_pnl"] / row["start_equity"] * 100 if row["start_equity"] else 0.0
        row["cum_pnl"] = cumulative
        row["cum_return_pct"] = cumulative / start_equity * 100 if start_equity else 0.0
        daily.append(row)

    by_coin = []
    for symbol in list(SYMBOLS) + sorted({t["symbol"] for t in trades} - set(SYMBOLS)):
        meta = coin_meta(symbol)
        rows = [t for t in trades if t["symbol"] == symbol]
        coin_net = sum(t["net_pnl"] or 0 for t in rows)
        by_coin.append({"symbol": symbol, "name": meta["name"], "ticker": meta["ticker"], "trades": len(rows),
                        "wins": sum(1 for t in rows if (t["net_pnl"] or 0) > 0), "net_pnl": coin_net,
                        "return_pct": coin_net / start_equity * 100 if start_equity else 0.0})

    # 차트용 흐름: 1d는 실시간 수익률 기록, 그 외는 일별 누적 수익률
    if period == "1d":
        points = storage.equity_series(mode, day=end_day)
        step = max(1, len(points) // 400)
        series = [{"t": p["ts"], "label": fmt_local(p["ts"], tz, "%H:%M"), "value": p["pnl_pct"]}
                  for p in points[::step]]
        if points and (len(points) - 1) % step:
            last = points[-1]
            series.append({"t": last["ts"], "label": fmt_local(last["ts"], tz, "%H:%M"), "value": last["pnl_pct"]})
    else:
        series = [{"t": d["day"], "label": d["day"][5:], "value": d["cum_return_pct"], "bar": d["return_pct"]}
                  for d in daily]

    # 하루(1d) 수익률은 대시보드과 같은 기준(미실현 손익 포함, 자산 변화 기준)으로 맞춘다
    realized_return = net / start_equity * 100 if start_equity else 0.0
    return_pct = series[-1]["value"] if period == "1d" and series else realized_return

    return {
        "period": period, "label": PERIODS[period][0], "mode": mode,
        "from": fmt_local(iso(start), tz, "%Y-%m-%d"), "to": fmt_local(iso(now), tz, "%Y-%m-%d"),
        "summary": {
            "trades": len(trades), "wins": len(wins), "losses": len(losses),
            "win_rate": len(wins) / len(trades) * 100 if trades else 0.0,
            "net_pnl": net, "gross_pnl": gross, "fee": fees, "start_equity": start_equity,
            "return_pct": return_pct, "realized_return_pct": realized_return,
            "best": max((t["net_pnl"] or 0 for t in trades), default=0.0),
            "worst": min((t["net_pnl"] or 0 for t in trades), default=0.0),
        },
        "series": series, "daily": daily, "by_coin": by_coin,
        "trades": [trade_view(t, tz) for t in reversed(trades)],
    }


# ---------------------------------------------------------------- 엑셀
HEAD_FILL = PatternFill("solid", fgColor="1F2A44")
HEAD_FONT = Font(bold=True, color="FFFFFF")
RED, BLUE = Font(color="C8322B"), Font(color="2457C5")   # 한국식: 수익 빨강, 손실 파랑


def _header(sheet, titles: list[str], row: int = 1) -> None:
    for col, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=col, value=title)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = f"A{row + 1}"   # 문자열로 지정해야 빈 행이 생기지 않는다


def _autosize(sheet, widths: dict[int, int] | None = None) -> None:
    for col in range(1, sheet.max_column + 1):
        longest = max((len(str(c.value)) for c in sheet[get_column_letter(col)] if c.value is not None), default=6)
        sheet.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, min(48, longest * 1.4 + 3))


def export_xlsx(storage, mode: str, period: str, tz: str, now=None) -> BytesIO:
    report = build_report(storage, mode, period, tz, now)
    s = report["summary"]
    book = Workbook()

    # 1) 요약
    sheet = book.active
    sheet.title = "요약"
    sheet["A1"] = f"{APP_NAME} 거래 리포트"
    sheet["A1"].font = Font(bold=True, size=14)
    rows = [
        ("기간", f"{report['label']} ({report['from']} ~ {report['to']})"),
        ("거래 모드", MODES.get(mode, mode)),
        ("시간대", tz),
        ("시작 자산 (USDT)", round(s["start_equity"], 4)),
        ("순손익 (USDT, 수수료 차감 후)", round(s["net_pnl"], 4)),
        ("수익률 (%)", round(s["return_pct"], 2)),
        ("총 거래 횟수", s["trades"]),
        ("수익 거래", s["wins"]),
        ("손실 거래", s["losses"]),
        ("승률 (%)", round(s["win_rate"], 1)),
        ("수수료 합계 (USDT)", round(s["fee"], 4)),
        ("수수료 차감 전 손익 (USDT)", round(s["gross_pnl"], 4)),
        ("가장 큰 수익 (USDT)", round(s["best"], 4)),
        ("가장 큰 손실 (USDT)", round(s["worst"], 4)),
    ]
    for index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=index, column=1, value=label).font = Font(bold=True)
        cell = sheet.cell(row=index, column=2, value=value)
        if label.startswith(("순손익", "수익률")) and isinstance(value, (int, float)):
            cell.font = RED if value > 0 else BLUE if value < 0 else Font()
    _autosize(sheet, {1: 34, 2: 34})

    # 2) 거래내역
    sheet = book.create_sheet("거래내역")
    _header(sheet, ["번호", "코인", "포지션", "레버리지", "수량", "진입 시각", "청산 시각", "진입가", "청산가",
                    "사용 증거금", "사용 원금 비율(%)", "수수료 전 손익", "수수료", "순손익",
                    "수익률(증거금 대비 %)", "수익률(자산 대비 %)", "청산 사유", "판단 주체", "진입 이유"])
    for t in reversed(report["trades"]):
        sheet.append([t["id"], t["ticker"], "롱" if t["side"] == "long" else "숏", t["leverage"], t["qty"],
                      t["entry_at"], t["exit_at"], t["entry_price"], t["exit_price"],
                      round(t["margin"] or 0, 4), round(t["invest_pct"] or 0, 2), round(t["gross_pnl"] or 0, 4),
                      round(t["fee"] or 0, 4), round(t["net_pnl"] or 0, 4), round(t["return_pct"] or 0, 2),
                      round(t["equity_return_pct"] or 0, 3), t["exit_reason"],
                      "Gemini" if t["source"] == "gemini" else "규칙 전략", t["entry_reason"]])
        value = t["net_pnl"] or 0
        for col in (14, 15, 16):
            sheet.cell(row=sheet.max_row, column=col).font = RED if value > 0 else BLUE if value < 0 else Font()
    _autosize(sheet, {19: 70})

    # 3) 일별수익 + 차트
    sheet = book.create_sheet("일별수익")
    _header(sheet, ["날짜", "거래 횟수", "순손익", "수수료", "일 수익률(%)", "누적 순손익", "누적 수익률(%)"])
    for d in report["daily"]:
        sheet.append([d["day"], d["trades"], round(d["net_pnl"], 4), round(d["fee"], 4),
                      round(d["return_pct"], 2), round(d["cum_pnl"], 4), round(d["cum_return_pct"], 2)])
    _autosize(sheet)
    count = len(report["daily"])
    if count:
        labels = Reference(sheet, min_col=1, min_row=2, max_row=count + 1)
        line = LineChart()
        line.title, line.height, line.width = "누적 수익률 (%)", 8, 18
        line.add_data(Reference(sheet, min_col=7, min_row=1, max_row=count + 1), titles_from_data=True)
        line.set_categories(labels)
        sheet.add_chart(line, "I2")
        bars = BarChart()
        bars.title, bars.height, bars.width = "일 수익률 (%)", 8, 18
        bars.add_data(Reference(sheet, min_col=5, min_row=1, max_row=count + 1), titles_from_data=True)
        bars.set_categories(labels)
        sheet.add_chart(bars, "I19")

    # 4) 코인별
    sheet = book.create_sheet("코인별")
    _header(sheet, ["코인", "거래 횟수", "수익 거래", "순손익", "수익률(자산 대비 %)"])
    for c in report["by_coin"]:
        sheet.append([f"{c['name']} ({c['ticker']})", c["trades"], c["wins"], round(c["net_pnl"], 4),
                      round(c["return_pct"], 2)])
    _autosize(sheet)

    # 5) 실시간 수익률 (하루 리포트일 때만)
    if period == "1d" and report["series"]:
        sheet = book.create_sheet("당일수익률")
        _header(sheet, ["시각", "당일 수익률(%)"])
        for p in report["series"]:
            sheet.append([p["label"], round(p["value"], 3)])
        chart = LineChart()
        chart.title, chart.height, chart.width = "당일 실시간 수익률 (%)", 8, 18
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=len(report["series"]) + 1), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=len(report["series"]) + 1))
        sheet.add_chart(chart, "D2")
        _autosize(sheet)

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer
